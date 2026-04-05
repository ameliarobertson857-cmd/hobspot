import argparse
import datetime as dt
import json
import os
import re
import shutil
import unicodedata
from pathlib import Path

import ai_classify_one as classifier_engine
from openpyxl import load_workbook
from playwright.sync_api import TimeoutError, sync_playwright

CLASSIFICATION_FILE = "classification_result.json"
CLASSIFICATION_TEXT_FILE = "classification_source_text.txt"
LEGACY_PDF_FILE = "download.pdf"
PDF_DIRECTORY = "downloaded_pdfs"
PROCESSED_PDFS_FILE = "vec_processed_pdfs.json"
PUBLIC_VEC_URL = "https://www.sspa.juntadeandalucia.es/servicioandaluzdesalud/profesionales/ventanilla-electronica-de-profesionales"
USER_DATA_DIR = "vec_browser_profile"
SCRIPT_VERSION = "2026-03-30-dashboard-merit-loop"
FINAL_SUBMIT = True
DEFAULT_TARGET_SUCCESSFUL_UPLOADS = 3
DEFAULT_CLASSIFICATION_BACKEND = "openai"
MAX_CLASSIFICATION_ATTEMPTS = 50

DEBUG_DIR = Path("vec_debug")
DEBUG_DIR.mkdir(exist_ok=True)
UPLOAD_TMP_DIR = Path("vec_upload_tmp")
UPLOAD_TMP_DIR.mkdir(exist_ok=True)
PROCESSED_PDF_ARCHIVE_DIR = Path("uploaded_pdfs")
PROCESSED_PDF_ARCHIVE_DIR.mkdir(exist_ok=True)
SKIPPED_CLASSIFICATION_DIR = Path("skipped_classifications")
SKIPPED_CLASSIFICATION_DIR.mkdir(exist_ok=True)
UPLOAD_DESCRIPTION_MAX_LENGTH = 100
PUBLIC_VEC_GOTO_TIMEOUT_MS = 120000
PUBLIC_VEC_RETRY_WAIT_MS = 3000
VEC_SESSION_READY_TIMEOUT_MS = 600000
TRACKING_WORKBOOK_ENV_VAR = "VEC_TRACKING_WORKBOOK"
TRACKING_WORKBOOK_DIR = Path.home() / "Downloads"
TRACKING_WORKBOOK_GLOB = "COPIA*UI PATH.xlsx"
TRACKING_MAIN_SHEET = "Estudio VEC"
TRACKING_LOG_SHEET = "VEC_UPLOAD_LOG"
TRACKING_REGULAR_ROW_START = 8
TRACKING_REGULAR_ROW_END = 60
TRACKING_UNIVERSITY_ROW_START = 62
TRACKING_UNIVERSITY_ROW_END = 86
PENDING_EXCEL_SYNC_FILE = Path("vec_pending_excel_sync.json")
TRACKING_LOG_HEADERS = [
    "document_key",
    "submitted_at",
    "sync_status",
    "excel_section",
    "excel_row",
    "source_pdf_path",
    "source_pdf_filename",
    "archived_pdf_path",
    "vec_route",
    "merit_name",
    "year",
    "hours_value",
    "cfc_credits_value",
    "ects_credits_value",
    "start_date",
    "end_date",
    "course_code",
    "professional_category",
    "institution_name",
    "center_result_text",
    "training_type",
    "delivery_method",
    "accredited_choice",
    "accredited_received",
    "document_description",
    "alerts",
    "confidence",
]

WIZARD_STEP_MARKERS = [
    "Fecha de Inicio",
    "Fecha Inicio",
    "Fecha Fin",
    "Fecha Fin/Obtención",
    "Fecha Fin/Obtenci",
    "Nombre del Mérito",
    "Nombre del Merito",
    "Centro",
    "Añadir documento",
    "Anadir documento",
    "Firmar",
]

STEP1_DATE_MARKERS = [
    "Fecha de Inicio",
    "Fecha Inicio",
    "Start date",
    "Fecha Fin",
    "Fecha Fin/Obtenci",
]

STEP2_CENTER_MARKERS = [
    "Centro",
    "Center",
]

STEP3_MERIT_MARKERS = [
    "Nombre del Mérito",
    "Nombre del Merito",
    "Name of Merit",
]

STEP4_SPECIFIC_MARKERS = [
    "Tipo Formación",
    "Tipo Formacion",
    "Ámbito",
    "Ambito",
    "Metodología",
    "Metodologia",
    "Metodología de Impartición",
    "Metodologia de Imparticion",
    "Tipo Título Propio/Diploma",
    "Tipo Titulo Propio/Diploma",
    "Créditos/Horas",
    "Creditos/Horas",
    "Horas",
    "Créditos",
    "Creditos",
    "Código curso",
    "Codigo curso",
]

STEP5_DOCUMENT_MARKERS = [
    "Añadir documento",
    "Anadir documento",
    "Subir archivo",
    "Subir fichero",
    "Adjuntar",
]

STEP6_SIGN_MARKERS = [
    "Firmar",
    "Firmar y registrar",
    "Registrar",
    "Enviar",
    "Submit",
]

STEP1_ACTIVE_TITLES = ["Fechas", "Dates"]
STEP2_ACTIVE_TITLES = ["Selección del Centro", "Seleccion del Centro", "Center Selection"]
STEP3_ACTIVE_TITLES = ["Nombre del Mérito", "Nombre del Merito", "Name of Merit"]
STEP4_ACTIVE_TITLES = ["Valores Específicos", "Valores Especificos", "Specific values"]
STEP5_ACTIVE_TITLES = ["Selección de los Documentos", "Seleccion de los Documentos", "Document Selection"]
STEP6_ACTIVE_TITLES = ["Firma de los Documentos", "Signing", "Document Signing"]

STEP4_GENERIC_MARKERS = [
    "Valores Específicos",
    "Valores Especificos",
    "Guardar y Siguiente",
    "Guardar y Salir",
]

PRIVATE_VEC_URL_TOKEN = "/VEC/faces/pages/private/"
PRIVATE_MERIT_URL_TOKEN = "/VEC/faces/pages/private/meritos/"
PUBLIC_VEC_LOGIN_URL_TOKEN = "/VEC/faces/pages/public/login/autoregistro.xhtml"
TRAINING_ROUTE_VALUES = {
    "formacion_continuada",
    "diplomas_titulos_propios_universitarios",
    "estancias_formativas",
}
BLOCKING_CLASSIFICATION_ALERT_MARKERS = (
    "posible documento solo de enlace o validacion",
    "clasificacion incompleta",
)
DETAIL_PAGE_MARKERS = ["Detalle", "Volver", "Editar", "Eliminar"]
POST_LOGIN_DASHBOARD_MARKERS = [
    "Acciones disponibles desde la pantalla principal",
    "Mis Datos",
    "Mensajes",
    "Ultimas Noticias",
    "Mis Procesos Selectivos",
    "Bolsa de Empleo",
]
NEW_MERIT_BUTTON_TEXTS = [
    "Nuevo Mérito",
    "Nuevo Merito",
    "Crear Nuevo Mérito",
    "Crear Nuevo Merito",
    "Añadir",
    "Anadir",
    "Nuevo",
    "Crear",
    "Alta",
]
MERIT_CATALOG_TREE_ROOT = ["Catalogo de Tipos de Meritos"]


ROUTE_LABELS = {
    "formacion_continuada": {
        "selected": ["Formación Continuada", "Formacion Continuada"],
        "tree_parent": ["Formación Continuada Recibida", "Formacion Continuada Recibida"],
        "tree_child": ["Formación Continuada", "Formacion Continuada"],
    },
    "diplomas_titulos_propios_universitarios": {
        "selected": [
            "Diplomas y Títulos Propios Universitarios",
            "Diplomas y Titulos Propios Universitarios",
        ],
        "tree_parent": ["Formación Continuada Recibida", "Formacion Continuada Recibida"],
        "tree_child": [
            "Diplomas y Títulos Propios Universitarios",
            "Diplomas y Titulos Propios Universitarios",
        ],
    },
    "estancias_formativas": {
        "selected": ["Estancias Formativas"],
        "tree_parent": ["Formación Continuada Recibida", "Formacion Continuada Recibida"],
        "tree_child": ["Estancias Formativas"],
    },
    "experiencia_sas": {
        "selected": ["Experiencia SAS"],
        "tree_parent": ["Experiencia Profesional"],
        "tree_child": ["Experiencia SAS"],
    },
}

ROUTE_TREE_PATHS = {
    "formacion_continuada": [
        MERIT_CATALOG_TREE_ROOT,
        ["FormaciÃ³n Continuada Recibida", "Formacion Continuada Recibida"],
        ["FormaciÃ³n Continuada", "Formacion Continuada"],
    ],
    "diplomas_titulos_propios_universitarios": [
        MERIT_CATALOG_TREE_ROOT,
        ["FormaciÃ³n Continuada Recibida", "Formacion Continuada Recibida"],
        ["FormaciÃ³n Continuada", "Formacion Continuada"],
        [
            "Diplomas y TÃ­tulos Propios Universitarios",
            "Diplomas y Titulos Propios Universitarios",
        ],
    ],
    "estancias_formativas": [
        MERIT_CATALOG_TREE_ROOT,
        ["FormaciÃ³n Continuada Recibida", "Formacion Continuada Recibida"],
        ["FormaciÃ³n Continuada", "Formacion Continuada"],
        ["Estancias Formativas"],
    ],
    "experiencia_sas": [
        MERIT_CATALOG_TREE_ROOT,
        ["Experiencia Profesional"],
        ["Experiencia SAS"],
    ],
}


def parse_args():
    parser = argparse.ArgumentParser(
        description="Fill the VEC merit wizard in batches using AI-classified PDFs."
    )
    parser.add_argument(
        "--target-successes",
        type=int,
        default=DEFAULT_TARGET_SUCCESSFUL_UPLOADS,
        help="Stop after this many documents have been uploaded successfully.",
    )
    parser.add_argument(
        "--classification-backend",
        choices=("openai", "heuristic"),
        default=DEFAULT_CLASSIFICATION_BACKEND,
        help="Classification backend used for each next PDF selection.",
    )
    return parser.parse_args()


def resolve_pdf_path(data):
    candidates = []
    seen = set()

    def add_candidate(path_value):
        if not path_value:
            return

        path = Path(path_value)
        normalized = os.path.normcase(os.path.normpath(str(path)))
        if normalized in seen:
            return

        seen.add(normalized)
        candidates.append(path)

    add_candidate(data.get("source_pdf_path"))

    source_pdf_filename = data.get("source_pdf_filename")
    add_candidate(source_pdf_filename)
    if source_pdf_filename:
        add_candidate(Path(PDF_DIRECTORY) / source_pdf_filename)

    add_candidate(LEGACY_PDF_FILE)

    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return str(candidate)

    searched_paths = ", ".join(str(candidate) for candidate in candidates) or "no candidate paths"
    raise FileNotFoundError(
        f"PDF file not found. Checked: {searched_paths}"
    )


def build_upload_pdf_path(pdf_path):
    source_path = Path(pdf_path)
    suffix = source_path.suffix or ".pdf"
    target_path = UPLOAD_TMP_DIR / f"vec_upload_document{suffix}"

    try:
        shutil.copy2(source_path, target_path)
    except Exception as exc:
        raise RuntimeError(f"Could not stage a short upload filename for {source_path.name!r}: {exc}") from exc

    return str(target_path)


def normalize_pdf_path_key(path_value):
    if not path_value:
        return ""

    try:
        resolved = Path(path_value).resolve(strict=False)
    except Exception:
        resolved = Path(path_value)

    return os.path.normcase(os.path.normpath(str(resolved)))


def load_processed_pdf_registry():
    registry_path = Path(PROCESSED_PDFS_FILE)
    if not registry_path.exists():
        return {"processed_pdfs": []}

    try:
        data = json.loads(registry_path.read_text(encoding="utf-8"))
    except Exception:
        return {"processed_pdfs": []}

    if isinstance(data, list):
        return {"processed_pdfs": data}
    if isinstance(data, dict) and isinstance(data.get("processed_pdfs"), list):
        return data
    return {"processed_pdfs": []}


def save_processed_pdf_registry(registry):
    registry_path = Path(PROCESSED_PDFS_FILE)
    registry_path.write_text(
        json.dumps(registry, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def record_processed_pdf(source_pdf_path, archived_pdf_path=""):
    source_path = Path(source_pdf_path) if source_pdf_path else None
    source_key = normalize_pdf_path_key(source_path)
    source_name = source_path.name if source_path else ""

    registry = load_processed_pdf_registry()
    records = registry.setdefault("processed_pdfs", [])
    timestamp = dt.datetime.now().isoformat(timespec="seconds")

    existing_record = None
    for record in records:
        record_source = normalize_pdf_path_key(record.get("source_pdf_path"))
        record_archived = normalize_pdf_path_key(record.get("archived_pdf_path"))
        if source_key and source_key in (record_source, record_archived):
            existing_record = record
            break
        if source_name and record.get("source_pdf_filename") == source_name:
            existing_record = record
            break

    payload = {
        "processed_at": timestamp,
        "source_pdf_path": str(source_path) if source_path else "",
        "source_pdf_filename": source_name,
        "archived_pdf_path": str(archived_pdf_path or ""),
    }

    if existing_record is None:
        records.append(payload)
    else:
        existing_record.update(payload)

    save_processed_pdf_registry(registry)


def next_archived_pdf_path(source_path):
    destination = PROCESSED_PDF_ARCHIVE_DIR / source_path.name
    if not destination.exists():
        return destination

    stem = source_path.stem
    suffix = source_path.suffix
    counter = 2
    while True:
        candidate = PROCESSED_PDF_ARCHIVE_DIR / f"{stem}-{counter}{suffix}"
        if not candidate.exists():
            return candidate
        counter += 1


def archive_uploaded_pdf(data):
    source_pdf_path = data.get("resolved_pdf_path") or data.get("source_pdf_path")
    if not source_pdf_path:
        return ""

    source_path = Path(source_pdf_path)
    if not source_path.exists():
        record_processed_pdf(source_pdf_path)
        return ""

    destination = next_archived_pdf_path(source_path)
    shutil.move(str(source_path), str(destination))
    record_processed_pdf(source_pdf_path, archived_pdf_path=str(destination))
    return str(destination)


def next_skipped_classification_path(source_pdf_path):
    source_name = Path(source_pdf_path or "classification").stem or "classification"
    destination = SKIPPED_CLASSIFICATION_DIR / f"{source_name}.json"
    if not destination.exists():
        return destination

    counter = 2
    while True:
        candidate = SKIPPED_CLASSIFICATION_DIR / f"{source_name}-{counter}.json"
        if not candidate.exists():
            return candidate
        counter += 1


def save_skipped_classification_snapshot(
    classification,
    data,
    blockers,
    reason,
    archived_pdf_path="",
):
    source_pdf_path = (
        classification.get("pdf_path")
        or data.get("resolved_pdf_path")
        or data.get("source_pdf_path")
        or ""
    )
    snapshot_path = next_skipped_classification_path(source_pdf_path)
    payload = {
        "skipped_at": dt.datetime.now().isoformat(timespec="seconds"),
        "source_pdf_path": source_pdf_path,
        "source_pdf_filename": data.get("source_pdf_filename") or Path(source_pdf_path or "").name,
        "archived_pdf_path": str(archived_pdf_path or ""),
        "reason": str(reason or "").strip(),
        "blockers": list(blockers or []),
        "classification_backend": classification.get("backend"),
        "skipped_candidates_before_selection": classification.get("skipped_candidates", []),
        "result": data,
    }
    snapshot_path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return str(snapshot_path)


def clean_excel_text(value, max_length=None):
    text = repair_mojibake_text(value)
    text = re.sub(r"\s+", " ", str(text or "")).strip()
    if max_length and len(text) > max_length:
        return text[:max_length].rstrip()
    return text


def safe_positive_number(value):
    if value in (None, ""):
        return None

    try:
        numeric_value = float(value)
    except (TypeError, ValueError):
        return None

    if numeric_value <= 0:
        return None
    if abs(numeric_value - round(numeric_value)) < 1e-9:
        return int(round(numeric_value))
    return round(numeric_value, 4)


def extract_tracking_year(data):
    for candidate in (
        data.get("year"),
        data.get("end_date"),
        data.get("start_date"),
    ):
        try:
            parsed_year = classifier_engine.parse_result_year(candidate)
        except Exception:
            parsed_year = None
        if parsed_year:
            return int(parsed_year)

        match = re.search(r"(19|20)\d{2}", str(candidate or ""))
        if match:
            return int(match.group(0))

    return None


def build_excel_document_key(data, archived_pdf_path=""):
    for candidate in (
        archived_pdf_path,
        data.get("resolved_pdf_path"),
        data.get("source_pdf_path"),
    ):
        key = normalize_pdf_path_key(candidate)
        if key:
            return key

    fallback_parts = [
        clean_excel_text(data.get("source_pdf_filename")),
        clean_excel_text(data.get("merit_name")),
        str(extract_tracking_year(data) or ""),
    ]
    fallback_text = "|".join(part for part in fallback_parts if part)
    return normalize_text(fallback_text)


def choose_tracking_section(vec_route):
    if vec_route == "diplomas_titulos_propios_universitarios":
        return "university_training", TRACKING_UNIVERSITY_ROW_START, TRACKING_UNIVERSITY_ROW_END
    if vec_route in TRAINING_ROUTE_VALUES:
        return "regular_training", TRACKING_REGULAR_ROW_START, TRACKING_REGULAR_ROW_END
    return "log_only", None, None


def build_excel_sync_entry(data, archived_pdf_path="", submitted_at=None):
    route = clean_excel_text(data.get("vec_route"))
    credits_or_hours = normalize_text(data.get("credits_or_hours"))
    credits_value = safe_positive_number(data.get("credits"))
    hours_value = safe_positive_number(data.get("hours"))
    sheet_section, row_start, row_end = choose_tracking_section(route)

    excel_hours_value = None
    excel_cfc_credits_value = None
    excel_ects_credits_value = None

    if route == "diplomas_titulos_propios_universitarios":
        if credits_value is not None:
            excel_ects_credits_value = credits_value
        elif hours_value is not None:
            excel_hours_value = hours_value
    else:
        if "credit" in credits_or_hours and credits_value is not None:
            excel_cfc_credits_value = credits_value
        elif "hora" in credits_or_hours and hours_value is not None:
            excel_hours_value = hours_value
        elif credits_value is not None and hours_value is None:
            excel_cfc_credits_value = credits_value
        elif hours_value is not None:
            excel_hours_value = hours_value

    source_pdf_filename = clean_excel_text(data.get("source_pdf_filename"))
    merit_name = clean_excel_text(
        data.get("merit_name")
        or data.get("document_description")
        or data.get("description")
        or Path(source_pdf_filename or "document").stem,
        max_length=250,
    )

    return {
        "document_key": build_excel_document_key(data, archived_pdf_path=archived_pdf_path),
        "submitted_at": submitted_at or dt.datetime.now().isoformat(timespec="seconds"),
        "excel_section": sheet_section,
        "row_start": row_start,
        "row_end": row_end,
        "source_pdf_path": clean_excel_text(data.get("resolved_pdf_path") or data.get("source_pdf_path")),
        "source_pdf_filename": source_pdf_filename,
        "archived_pdf_path": clean_excel_text(archived_pdf_path),
        "vec_route": route,
        "merit_name": merit_name,
        "year": extract_tracking_year(data),
        "hours_value": excel_hours_value,
        "cfc_credits_value": excel_cfc_credits_value,
        "ects_credits_value": excel_ects_credits_value,
        "start_date": clean_excel_text(data.get("start_date")),
        "end_date": clean_excel_text(data.get("end_date")),
        "course_code": clean_excel_text(data.get("course_code")),
        "professional_category": clean_excel_text(data.get("professional_category")),
        "institution_name": clean_excel_text(data.get("institution_name")),
        "center_result_text": clean_excel_text(data.get("center_result_text")),
        "training_type": clean_excel_text(data.get("training_type")),
        "delivery_method": clean_excel_text(data.get("delivery_method")),
        "accredited_choice": clean_excel_text(data.get("accredited_choice")),
        "accredited_received": clean_excel_text(data.get("accredited_received")),
        "document_description": clean_excel_text(data.get("document_description")),
        "alerts": " | ".join(clean_excel_text(item) for item in (data.get("alerts") or []) if clean_excel_text(item)),
        "confidence": data.get("confidence"),
    }


def load_pending_excel_syncs():
    if not PENDING_EXCEL_SYNC_FILE.exists():
        return []

    try:
        payload = json.loads(PENDING_EXCEL_SYNC_FILE.read_text(encoding="utf-8"))
    except Exception:
        return []

    if isinstance(payload, list):
        return payload
    return []


def save_pending_excel_syncs(entries):
    if entries:
        PENDING_EXCEL_SYNC_FILE.write_text(
            json.dumps(entries, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
    elif PENDING_EXCEL_SYNC_FILE.exists():
        PENDING_EXCEL_SYNC_FILE.unlink()


def queue_pending_excel_sync(entry, error_message):
    queued_entries = load_pending_excel_syncs()
    updated = False

    for queued_entry in queued_entries:
        if queued_entry.get("document_key") == entry.get("document_key"):
            queued_entry.update(entry)
            queued_entry["sync_error"] = clean_excel_text(error_message, max_length=500)
            updated = True
            break

    if not updated:
        queued_payload = dict(entry)
        queued_payload["sync_error"] = clean_excel_text(error_message, max_length=500)
        queued_entries.append(queued_payload)

    save_pending_excel_syncs(queued_entries)


def resolve_tracking_workbook_path():
    configured_path = clean_excel_text(os.getenv(TRACKING_WORKBOOK_ENV_VAR))
    if configured_path:
        workbook_path = Path(configured_path).expanduser()
        if workbook_path.exists():
            return workbook_path
        raise FileNotFoundError(
            f"Tracking workbook set in {TRACKING_WORKBOOK_ENV_VAR} was not found: {workbook_path}"
        )

    matches = sorted(
        TRACKING_WORKBOOK_DIR.glob(TRACKING_WORKBOOK_GLOB),
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )
    if matches:
        return matches[0]

    raise FileNotFoundError(
        f"Could not find a workbook matching {TRACKING_WORKBOOK_GLOB!r} in {TRACKING_WORKBOOK_DIR}"
    )


def ensure_tracking_log_sheet(workbook):
    if TRACKING_LOG_SHEET in workbook.sheetnames:
        log_ws = workbook[TRACKING_LOG_SHEET]
    else:
        log_ws = workbook.create_sheet(TRACKING_LOG_SHEET)

    current_headers = [log_ws.cell(1, column).value for column in range(1, len(TRACKING_LOG_HEADERS) + 1)]
    if current_headers != TRACKING_LOG_HEADERS:
        for index, header in enumerate(TRACKING_LOG_HEADERS, start=1):
            log_ws.cell(1, index).value = header

    return log_ws


def find_existing_log_row(log_ws, entry):
    key_column = TRACKING_LOG_HEADERS.index("document_key") + 1
    source_name_column = TRACKING_LOG_HEADERS.index("source_pdf_filename") + 1
    document_key = clean_excel_text(entry.get("document_key"))
    source_pdf_filename = clean_excel_text(entry.get("source_pdf_filename"))

    for row_number in range(2, log_ws.max_row + 1):
        row_key = clean_excel_text(log_ws.cell(row_number, key_column).value)
        row_source_name = clean_excel_text(log_ws.cell(row_number, source_name_column).value)
        if document_key and row_key == document_key:
            return row_number
        if source_pdf_filename and row_source_name == source_pdf_filename:
            return row_number

    return None


def find_next_empty_tracking_row(ws, row_start, row_end):
    for row_number in range(row_start, row_end + 1):
        if all(ws.cell(row_number, column).value in (None, "") for column in range(1, 6)):
            return row_number
    return None


def write_tracking_input_row(ws, row_number, entry):
    ws.cell(row_number, 1).value = entry["merit_name"]
    ws.cell(row_number, 2).value = entry["year"]
    ws.cell(row_number, 3).value = entry["hours_value"]
    ws.cell(row_number, 4).value = entry["cfc_credits_value"]
    ws.cell(row_number, 5).value = entry["ects_credits_value"]


def append_tracking_log_row(log_ws, entry, sync_status, excel_row=None):
    target_row = log_ws.max_row + 1
    payload = dict(entry)
    payload["sync_status"] = sync_status
    payload["excel_row"] = excel_row

    for index, header in enumerate(TRACKING_LOG_HEADERS, start=1):
        value = payload.get(header)
        log_ws.cell(target_row, index).value = value


def sync_excel_entry(entry):
    workbook_path = resolve_tracking_workbook_path()
    workbook = load_workbook(workbook_path)

    if TRACKING_MAIN_SHEET not in workbook.sheetnames:
        raise KeyError(f"Worksheet {TRACKING_MAIN_SHEET!r} was not found in {workbook_path}")

    input_ws = workbook[TRACKING_MAIN_SHEET]
    log_ws = ensure_tracking_log_sheet(workbook)

    existing_log_row = find_existing_log_row(log_ws, entry)
    if existing_log_row is not None:
        return {
            "status": "already_recorded",
            "workbook_path": str(workbook_path),
            "excel_row": log_ws.cell(existing_log_row, TRACKING_LOG_HEADERS.index("excel_row") + 1).value,
            "log_row": existing_log_row,
        }

    excel_row = None
    status = "logged_only"
    if entry["excel_section"] != "log_only":
        excel_row = find_next_empty_tracking_row(input_ws, entry["row_start"], entry["row_end"])
        if excel_row is None:
            raise RuntimeError(
                f"No empty row is available in {TRACKING_MAIN_SHEET} for section {entry['excel_section']}"
            )
        write_tracking_input_row(input_ws, excel_row, entry)
        status = "synced"

    append_tracking_log_row(log_ws, entry, sync_status=status, excel_row=excel_row)
    workbook.save(workbook_path)

    return {
        "status": status,
        "workbook_path": str(workbook_path),
        "excel_row": excel_row,
        "log_row": log_ws.max_row,
    }


def sync_successful_upload_to_excel(data, archived_pdf_path=""):
    entry = build_excel_sync_entry(data, archived_pdf_path=archived_pdf_path)
    try:
        result = sync_excel_entry(entry)
    except Exception as exc:
        queue_pending_excel_sync(entry, str(exc))
        return {
            "status": "queued",
            "workbook_path": "",
            "excel_row": None,
            "log_row": None,
            "error": str(exc),
        }

    pending_entries = [
        item for item in load_pending_excel_syncs()
        if item.get("document_key") != entry.get("document_key")
    ]
    save_pending_excel_syncs(pending_entries)
    return result


def flush_pending_excel_syncs():
    queued_entries = load_pending_excel_syncs()
    if not queued_entries:
        return {"synced": 0, "remaining": 0}

    remaining_entries = []
    synced_count = 0

    for entry in queued_entries:
        try:
            sync_excel_entry(entry)
            synced_count += 1
        except Exception as exc:
            entry["sync_error"] = clean_excel_text(str(exc), max_length=500)
            remaining_entries.append(entry)

    save_pending_excel_syncs(remaining_entries)
    return {"synced": synced_count, "remaining": len(remaining_entries)}


def load_data():
    if not os.path.exists(CLASSIFICATION_FILE):
        raise FileNotFoundError(f"{CLASSIFICATION_FILE} not found")

    with open(CLASSIFICATION_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)

    data["resolved_pdf_path"] = resolve_pdf_path(data)

    return data


def save_debug(page, name):
    try:
        page.screenshot(path=str(DEBUG_DIR / f"{name}.png"), full_page=True)
    except Exception:
        pass

    try:
        (DEBUG_DIR / f"{name}.html").write_text(page.content(), encoding="utf-8")
    except Exception:
        pass

    try:
        frame_urls = [frame.url for frame in page.frames]
        debug_meta = {
            "url": page.url,
            "title": page.title(),
            "frame_urls": frame_urls,
        }
        (DEBUG_DIR / f"{name}.json").write_text(
            json.dumps(debug_meta, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
    except Exception:
        pass


def iter_scopes(page):
    if not page_is_usable(page):
        return

    yield page

    try:
        main_frame = page.main_frame
        frames = list(page.frames)
    except Exception:
        return

    for frame in frames:
        if frame == main_frame:
            continue
        yield frame


def first_visible(locator_list):
    for locator in locator_list:
        try:
            if locator.count() > 0 and locator.first.is_visible():
                return locator.first
        except Exception:
            continue
    return None


def first_existing(locator_list):
    for locator in locator_list:
        try:
            if locator.count() > 0:
                return locator.first
        except Exception:
            continue
    return None


def field_locators(page, labels):
    locators = []
    for label in expand_ui_text_variants(labels):
        locators.extend(
            [
                page.get_by_label(label, exact=False),
                page.locator(
                    f"xpath=(//label[contains(normalize-space(.), \"{label}\")]/following::*[self::input or self::textarea][not(@type='hidden')][1])[1]"
                ),
                page.locator(
                    f"xpath=(//*[self::span or self::div or self::td or self::th][contains(normalize-space(.), \"{label}\")]/following::*[self::input or self::textarea][not(@type='hidden')][1])[1]"
                ),
                page.locator(
                    f"xpath=(//*[contains(normalize-space(.), \"{label}\")]/ancestor::*[self::div or self::td or self::th or self::tr][1]//*[self::input or self::textarea][not(@type='hidden')][1])[1]"
                ),
                page.locator(
                    f"input[aria-label*=\"{label}\" i], textarea[aria-label*=\"{label}\" i], input[placeholder*=\"{label}\" i], textarea[placeholder*=\"{label}\" i]"
                ),
            ]
        )
    return locators


def scope_contains_markers(scope, texts):
    expanded_texts = expand_ui_text_variants(texts)

    for text in expanded_texts:
        try:
            if scope.get_by_text(text, exact=False).count() > 0:
                return True
        except Exception:
            continue

    scope_text = ""
    try:
        scope_text = scope.locator("body").inner_text(timeout=1500)
    except Exception:
        try:
            scope_text = scope.evaluate("() => document.body ? document.body.innerText : ''")
        except Exception:
            scope_text = ""

    normalized_scope_text = normalize_text(scope_text)
    if not normalized_scope_text:
        return False

    for text in expanded_texts:
        normalized_text = normalize_text(text)
        if normalized_text and normalized_text in normalized_scope_text:
            return True
    return False


def text_match_score(value, candidates):
    normalized_value = normalize_text(value)
    best_score = 0

    for index, candidate in enumerate(candidates):
        normalized_candidate = normalize_text(candidate)
        if not normalized_candidate:
            continue
        if normalized_value == normalized_candidate:
            best_score = max(best_score, 1000 - index)
        elif normalized_candidate in normalized_value:
            best_score = max(best_score, 800 - index)
        elif normalized_value in normalized_candidate:
            best_score = max(best_score, 600 - index)

    return best_score


def dedupe_preserve_order(values):
    seen = set()
    result = []
    for value in values:
        normalized = normalize_text(value)
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        result.append(value)
    return result


def safe_page_title(page):
    try:
        return page.title()
    except Exception:
        return "(title unavailable)"


def safe_page_url(page):
    try:
        return page.url
    except Exception:
        return "(url unavailable)"


def page_is_usable(page):
    try:
        return page is not None and not page.is_closed()
    except Exception:
        return False


def open_context_pages(context):
    try:
        pages = list(context.pages)
    except Exception:
        return []

    return [page for page in pages if page_is_usable(page)]


def print_open_pages(context):
    print("Open browser pages:")
    pages = open_context_pages(context)
    if not pages:
        print("  (no open pages)")
        return

    for index, candidate in enumerate(pages, start=1):
        print(f"  [{index}] {safe_page_url(candidate)} | {safe_page_title(candidate)}")


def save_context_debug(context, prefix):
    safe_prefix = re.sub(r"[^a-zA-Z0-9_-]+", "_", prefix).strip("_") or "debug"
    for index, candidate in enumerate(open_context_pages(context), start=1):
        save_debug(candidate, f"{safe_prefix}_page_{index}")


def safe_close_context(context):
    try:
        context.close()
        return True
    except Exception:
        return False


def bring_page_to_front(page):
    if not page_is_usable(page):
        return page

    try:
        page.bring_to_front()
    except Exception:
        pass
    return page


def is_private_vec_page(page):
    return PRIVATE_VEC_URL_TOKEN in (safe_page_url(page) or "")


def is_plain_public_vec_start_page(page):
    url = safe_page_url(page) or ""
    if PUBLIC_VEC_LOGIN_URL_TOKEN not in url and PUBLIC_VEC_URL not in url:
        return False

    try:
        if dashboard_ready(page) or merit_listing_ready(page) or merit_detail_page_ready(page):
            return False
    except Exception:
        pass

    return True


def choose_start_page(context):
    pages = open_context_pages(context)
    if not pages:
        return context.new_page()

    newest_page = pages[-1]
    if not is_plain_public_vec_start_page(newest_page):
        return bring_page_to_front(newest_page)

    for candidate in reversed(pages):
        if is_private_merit_page(candidate):
            return bring_page_to_front(candidate)

    for candidate in reversed(pages):
        if is_private_vec_page(candidate):
            return bring_page_to_front(candidate)

    for candidate in reversed(pages):
        if not is_plain_public_vec_start_page(candidate):
            return bring_page_to_front(candidate)

    return bring_page_to_front(newest_page)


def open_public_vec_start_page(context):
    page = choose_start_page(context)
    navigation_errors = []

    for attempt in range(1, 4):
        for wait_until in ("domcontentloaded", "commit"):
            try:
                print(
                    f"Opening VEC public page (attempt {attempt}/3, wait_until={wait_until}, "
                    f"timeout={PUBLIC_VEC_GOTO_TIMEOUT_MS}ms)..."
                )
                page.goto(PUBLIC_VEC_URL, wait_until=wait_until, timeout=PUBLIC_VEC_GOTO_TIMEOUT_MS)
                page.wait_for_timeout(2000)
                return page
            except TimeoutError as exc:
                navigation_errors.append(f"attempt {attempt} {wait_until}: timeout")
                print(f"VEC public page load is taking longer than expected: {exc}")
            except Exception as exc:
                navigation_errors.append(f"attempt {attempt} {wait_until}: {exc}")
                print(f"Could not auto-open the VEC public page: {exc}")

            current_url = safe_page_url(page) or ""
            if any(token in current_url for token in (PRIVATE_MERIT_URL_TOKEN, "juntadeandalucia.es")):
                return page

            if attempt < 3:
                page.wait_for_timeout(PUBLIC_VEC_RETRY_WAIT_MS)

    print()
    print("Automatic opening of the public VEC page failed or timed out.")
    print("You can still continue in this same browser window:")
    print("1. Check your internet connection")
    print("2. Open the VEC page manually")
    print("3. Complete access/login if needed")
    print("4. Open the exact wizard page")
    if navigation_errors:
        print("Startup navigation attempts:")
        for error in navigation_errors[-6:]:
            print(f"- {error}")
    return page


def repair_mojibake_text(value):
    text = str(value or "")
    mojibake_markers = ("Ã", "Â", "â", "€", "™", "œ", "ž", "ƒ")

    for _ in range(3):
        if not any(marker in text for marker in mojibake_markers):
            break

        try:
            repaired = text.encode("latin-1").decode("utf-8")
        except Exception:
            break

        if repaired == text:
            break
        text = repaired

    return text


def normalize_text(value):
    normalized = unicodedata.normalize("NFKD", repair_mojibake_text(value))
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", ascii_text).strip().casefold()


def expand_ui_text_variants(texts):
    alias_groups = {
        "buscar": ["Buscar", "Search"],
        "search": ["Buscar", "Search"],
        "centro": ["Centro", "Center"],
        "center": ["Centro", "Center"],
        "fecha de inicio": ["Fecha de Inicio", "Fecha Inicio", "Start date"],
        "fecha inicio": ["Fecha de Inicio", "Fecha Inicio", "Start date"],
        "start date": ["Fecha de Inicio", "Fecha Inicio", "Start date"],
        "fecha fin": [
            "Fecha Fin",
            "Fecha de Fin",
            "Fecha Fin/Obtencion",
            "Fecha Fin/Obtención",
            "Fecha de Obtencion",
            "Fecha de Obtención",
            "End date",
            "End/obtaining Date",
        ],
        "fecha fin/obtencion": [
            "Fecha Fin",
            "Fecha de Fin",
            "Fecha Fin/Obtencion",
            "Fecha Fin/Obtención",
            "Fecha de Obtencion",
            "Fecha de Obtención",
            "End date",
            "End/obtaining Date",
        ],
        "fecha de obtencion": [
            "Fecha Fin",
            "Fecha de Fin",
            "Fecha Fin/Obtencion",
            "Fecha Fin/Obtención",
            "Fecha de Obtencion",
            "Fecha de Obtención",
            "End date",
            "End/obtaining Date",
        ],
        "end date": [
            "Fecha Fin",
            "Fecha de Fin",
            "Fecha Fin/Obtencion",
            "Fecha Fin/Obtención",
            "Fecha de Obtencion",
            "Fecha de Obtención",
            "End date",
            "End/obtaining Date",
        ],
        "nombre del merito": ["Nombre del Merito", "Nombre del Mérito", "Name of Merit"],
        "name of merit": ["Nombre del Merito", "Nombre del Mérito", "Name of Merit"],
        "descripcion": ["Descripcion", "Descripción", "Description"],
        "description": ["Descripcion", "Descripción", "Description"],
        "descripcion adicional": ["Descripcion Adicional", "Descripción Adicional", "Additional Description"],
        "additional description": ["Descripcion Adicional", "Descripción Adicional", "Additional Description"],
        "ambito": ["Ambito", "Ámbito", "Scope"],
        "scope": ["Ambito", "Ámbito", "Scope"],
        "tipo formacion": ["Tipo Formacion", "Tipo Formación", "Training type"],
        "training type": ["Tipo Formacion", "Tipo Formación", "Training type"],
        "metodologia": ["Metodologia", "Metodología", "Methodology", "Delivery method"],
        "methodology": ["Metodologia", "Metodología", "Methodology", "Delivery method"],
        "delivery method": ["Metodologia", "Metodología", "Methodology", "Delivery method"],
        "acreditada": ["Acreditada", "Acreditada Recibida", "Accredited", "Accredited received"],
        "acreditada recibida": ["Acreditada", "Acreditada Recibida", "Accredited", "Accredited received"],
        "accredited": ["Acreditada", "Acreditada Recibida", "Accredited", "Accredited received"],
        "accredited received": ["Acreditada", "Acreditada Recibida", "Accredited", "Accredited received"],
        "codigo curso": ["Codigo curso", "Código curso", "Course code"],
        "course code": ["Codigo curso", "Código curso", "Course code"],
        "creditos": ["Creditos", "Créditos", "Credits"],
        "credits": ["Creditos", "Créditos", "Credits"],
        "horas": ["Horas", "Hours"],
        "hours": ["Horas", "Hours"],
        "tipo titulo propio/diploma": [
            "Tipo Titulo Propio/Diploma",
            "Tipo Título Propio/Diploma",
            "Tipo Titulo Propio",
            "Tipo Título Propio",
        ],
        "guardar y siguiente": ["Guardar y Siguiente", "Save and Next", "Siguiente", "Next", "Continuar", "Continue"],
        "save and next": ["Guardar y Siguiente", "Save and Next", "Siguiente", "Next", "Continuar", "Continue"],
        "siguiente": ["Guardar y Siguiente", "Save and Next", "Siguiente", "Next", "Continuar", "Continue"],
        "next": ["Guardar y Siguiente", "Save and Next", "Siguiente", "Next", "Continuar", "Continue"],
        "continuar": ["Guardar y Siguiente", "Save and Next", "Siguiente", "Next", "Continuar", "Continue"],
        "continue": ["Guardar y Siguiente", "Save and Next", "Siguiente", "Next", "Continuar", "Continue"],
        "aceptar": ["Aceptar", "Accept", "Confirmar", "Confirm"],
        "accept": ["Aceptar", "Accept", "Confirmar", "Confirm"],
        "confirmar": ["Aceptar", "Accept", "Confirmar", "Confirm"],
        "confirm": ["Aceptar", "Accept", "Confirmar", "Confirm"],
        "firmar documentos": ["Firmar Documentos", "Sign Documents"],
        "sign documents": ["Firmar Documentos", "Sign Documents"],
        "firmar y registrar": ["Firmar y registrar", "Sign and register", "Sign and submit"],
        "sign and register": ["Firmar y registrar", "Sign and register", "Sign and submit"],
        "firmar con servidor": ["Firmar con Servidor", "Firma con Servidor", "Sign with Server"],
        "sign with server": ["Firmar con Servidor", "Firma con Servidor", "Sign with Server"],
    }

    expanded = []
    seen = set()
    for text in texts or []:
        if text in (None, ""):
            continue

        raw_text = str(text)
        repaired_text = repair_mojibake_text(raw_text)
        variants = [raw_text, repaired_text]
        variants.extend(alias_groups.get(normalize_text(repaired_text), []))

        for variant in variants:
            normalized_variant = normalize_text(variant)
            if not normalized_variant or normalized_variant in seen:
                continue
            seen.add(normalized_variant)
            expanded.append(variant)

    return expanded


def text_matches(value, candidates):
    normalized_value = normalize_text(value)
    return any(normalize_text(candidate) == normalized_value for candidate in candidates if candidate)


def is_private_merit_page(page):
    return PRIVATE_MERIT_URL_TOKEN in (page.url or "")


def current_route_label(page):
    for scope in iter_scopes(page):
        try:
            selected = scope.locator(".ui-treenode-label.ui-state-highlight")
            if selected.count() > 0:
                return selected.first.inner_text().strip()
        except Exception:
            continue

    for scope in iter_scopes(page):
        try:
            heading = scope.locator("span.h1")
            if heading.count() > 0:
                return heading.first.inner_text().strip()
        except Exception:
            continue

    return ""


def current_active_step_title(page):
    for scope in iter_scopes(page):
        try:
            active_step = scope.locator(".ui-steps-item.ui-state-highlight .ui-steps-title")
            if active_step.count() > 0:
                return active_step.first.inner_text().strip()
        except Exception:
            continue

    return ""


def safe_wait_for_timeout(page, timeout_ms):
    if not page_is_usable(page):
        return False

    try:
        page.wait_for_timeout(timeout_ms)
        return True
    except Exception:
        return False


def wait_for_active_step(page, titles, timeout_ms=15000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        active_title = current_active_step_title(page)
        if active_title and text_matches(active_title, titles):
            return True
        if not safe_wait_for_timeout(page, 300):
            return False

    active_title = current_active_step_title(page)
    return bool(active_title and text_matches(active_title, titles))


def visible_error_messages(page):
    messages = []
    locators = [
        ".ui-growl-message p",
        ".ui-messages-error-summary",
        ".ui-messages-error-detail",
        ".ui-message-error-detail",
        ".ui-message-error-summary",
    ]

    for scope in iter_scopes(page):
        for selector in locators:
            try:
                candidates = scope.locator(selector)
                count = candidates.count()
            except Exception:
                continue

            for index in range(count):
                candidate = candidates.nth(index)
                try:
                    if not candidate.is_visible():
                        continue
                    text = candidate.inner_text().strip()
                except Exception:
                    continue

                if text:
                    messages.append(text)

    return dedupe_preserve_order(messages)


def merit_tree_scopes(page_or_scope):
    scopes = []

    try:
        tree_candidates = [
            page_or_scope.locator("#formMenuNavegacion\\:tipMerTree"),
            page_or_scope.locator("[id$='tipMerTree']"),
            page_or_scope.locator(".ui-tree"),
        ]
    except Exception:
        tree_candidates = []

    for locator in tree_candidates:
        try:
            count = locator.count()
        except Exception:
            count = 0

        for index in range(count):
            candidate = locator.nth(index)
            try:
                if candidate.is_visible():
                    scopes.append(candidate)
            except Exception:
                continue

    return scopes or [page_or_scope]


def find_tree_label(scope, labels, visible_only=True):
    for tree_scope in merit_tree_scopes(scope):
        try:
            candidates = tree_scope.locator(".ui-treenode-label")
            count = candidates.count()
        except Exception:
            continue

        for index in range(count):
            candidate = candidates.nth(index)
            try:
                if visible_only and not candidate.is_visible():
                    continue
                text = candidate.inner_text().strip()
            except Exception:
                continue

            if text_matches(text, labels):
                return candidate

    return None


def tree_node_content_from_label(label):
    if not label:
        return None

    try:
        return label.locator("xpath=ancestor::*[contains(@class, 'ui-treenode-content')][1]")
    except Exception:
        return None


def tree_children_scope_from_content(content):
    if not content:
        return None

    try:
        children = content.locator("xpath=following-sibling::*[contains(@class, 'ui-treenode-children')][1]")
        if children.count() > 0:
            return children.first
    except Exception:
        return None

    return None


def find_tree_label_under_parent(parent_content, labels, visible_only=True):
    child_scope = tree_children_scope_from_content(parent_content)
    if not child_scope:
        return None
    return find_tree_label(child_scope, labels, visible_only=visible_only)


def route_is_selected(page, labels):
    current_label = current_route_label(page)
    return bool(current_label and text_matches(current_label, labels))


def wait_for_route_selection(page, labels, timeout_ms=10000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        if route_is_selected(page, labels):
            return True
        if not safe_wait_for_timeout(page, 300):
            return False

    return route_is_selected(page, labels)


def wait_for_tree_branch_expanded(content, timeout_ms=5000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        try:
            if content.get_attribute("aria-expanded") != "false":
                return True
        except Exception:
            return True

        time.sleep(0.2)

    try:
        return content.get_attribute("aria-expanded") != "false"
    except Exception:
        return True


def wait_for_tree_label(page, labels, timeout_ms=5000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        for scope in iter_scopes(page):
            label = find_tree_label(scope, labels, visible_only=True)
            if label:
                return label
        if not safe_wait_for_timeout(page, 250):
            return None

    for scope in iter_scopes(page):
        label = find_tree_label(scope, labels, visible_only=True)
        if label:
            return label
    return None


def wait_for_tree_label_with_parent(page, labels, parent_content=None, timeout_ms=5000):
    if not parent_content:
        return wait_for_tree_label(page, labels, timeout_ms=timeout_ms)

    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        label = find_tree_label_under_parent(parent_content, labels, visible_only=True)
        if label:
            return label
        if not safe_wait_for_timeout(page, 250):
            return None

    return find_tree_label_under_parent(parent_content, labels, visible_only=True)


def expand_tree_branch(page, labels, parent_content=None):
    search_scopes = [parent_content] if parent_content else list(iter_scopes(page))
    for scope in search_scopes:
        branch = wait_for_tree_label_with_parent(page, labels, parent_content=parent_content, timeout_ms=1500)
        if not branch and not parent_content:
            branch = find_tree_label(scope, labels, visible_only=True)
        if not branch:
            continue

        content = tree_node_content_from_label(branch)
        if not content:
            continue
        try:
            expanded = content.get_attribute("aria-expanded")
        except Exception:
            expanded = None

        if expanded == "false":
            toggler = content.locator("xpath=./*[contains(@class, 'ui-tree-toggler')]")
            try:
                if toggler.count() > 0:
                    try:
                        toggler.first.scroll_into_view_if_needed(timeout=5000)
                    except Exception:
                        pass
                    toggler.first.click(timeout=5000)
                    page.wait_for_timeout(800)
                    if not wait_for_tree_branch_expanded(content, timeout_ms=5000):
                        return False
            except Exception:
                return False

        return True

    return False


def click_tree_leaf(page, labels, force=False, parent_content=None):
    search_scopes = [parent_content] if parent_content else list(iter_scopes(page))
    for scope in search_scopes:
        leaf = (
            find_tree_label_under_parent(parent_content, labels, visible_only=True)
            if parent_content
            else find_tree_label(scope, labels, visible_only=True)
        )
        if not leaf:
            continue

        content = tree_node_content_from_label(leaf)
        if not content:
            continue
        try:
            if content.get_attribute("aria-selected") == "true" and not force:
                return True
        except Exception:
            pass

        try:
            leaf.scroll_into_view_if_needed()
        except Exception:
            pass

        for candidate in (leaf, content):
            try:
                candidate.click(timeout=5000)
                page.wait_for_timeout(1000)
                if wait_for_route_selection(page, labels, timeout_ms=6000):
                    return True
            except Exception:
                continue

    return False


def expand_tree_until_label(page, target_labels, root_labels=None, timeout_ms=15000):
    import time

    if root_labels:
        expand_tree_branch(page, root_labels)

    if wait_for_tree_label(page, target_labels, timeout_ms=1200):
        return True

    end = time.time() + (timeout_ms / 1000)
    expanded_node_ids = set()

    while time.time() < end:
        expanded_any = False

        for scope in iter_scopes(page):
            for tree_scope in merit_tree_scopes(scope):
                try:
                    parent_nodes = tree_scope.locator("li.ui-treenode-parent")
                    count = parent_nodes.count()
                except Exception:
                    continue

                for index in range(count):
                    node = parent_nodes.nth(index)
                    try:
                        if not node.is_visible():
                            continue
                    except Exception:
                        continue

                    try:
                        node_id = node.get_attribute("id") or f"parent-{index}"
                    except Exception:
                        node_id = f"parent-{index}"

                    if node_id in expanded_node_ids:
                        continue

                    content = node.locator("xpath=./span[contains(@class, 'ui-treenode-content')]")
                    try:
                        if content.get_attribute("aria-expanded") != "false":
                            expanded_node_ids.add(node_id)
                            continue
                    except Exception:
                        expanded_node_ids.add(node_id)
                        continue

                    toggler = node.locator(".ui-tree-toggler")
                    try:
                        if toggler.count() == 0:
                            expanded_node_ids.add(node_id)
                            continue
                        toggler.first.scroll_into_view_if_needed(timeout=5000)
                        toggler.first.click(timeout=5000)
                        page.wait_for_timeout(700)
                    except Exception:
                        continue

                    expanded_node_ids.add(node_id)
                    expanded_any = True

                    if wait_for_tree_label(page, target_labels, timeout_ms=1200):
                        return True

        if not expanded_any:
            break

    return bool(wait_for_tree_label(page, target_labels, timeout_ms=1500))


def route_tree_segments(route_key, route_config):
    configured_path = ROUTE_TREE_PATHS.get(route_key)
    if configured_path:
        return configured_path

    segments = []
    if route_config.get("tree_parent"):
        segments.append(route_config["tree_parent"])
    if route_config.get("tree_child"):
        segments.append(route_config["tree_child"])
    return segments


def validate_route_tree_configuration():
    problems = []

    for route_key, route_config in ROUTE_LABELS.items():
        segments = route_tree_segments(route_key, route_config)
        if not segments:
            problems.append(f"{route_key}: missing route tree path")
            continue

        if not route_config.get("selected"):
            problems.append(f"{route_key}: missing selected labels")
        if not route_config.get("tree_child"):
            problems.append(f"{route_key}: missing tree child labels")

        root_segment = segments[0]
        if not text_matches(root_segment[0], MERIT_CATALOG_TREE_ROOT):
            problems.append(f"{route_key}: route tree path must start at the merit catalog root")

        leaf_segment = segments[-1]
        if route_config.get("tree_child") and not text_matches(leaf_segment[0], route_config["tree_child"]):
            problems.append(f"{route_key}: last route tree path segment must match the configured merit leaf")

    if problems:
        raise RuntimeError("Invalid merit route configuration: " + " | ".join(problems))


def open_route_tree_path(page, path_segments):
    if not path_segments:
        return True

    root_labels = path_segments[0]
    parent_content = None

    for labels in path_segments[:-1]:
        branch = wait_for_tree_label_with_parent(page, labels, parent_content=parent_content, timeout_ms=5000)
        if not branch:
            if not expand_tree_until_label(page, path_segments[-1], root_labels=root_labels, timeout_ms=15000):
                return False
            branch = wait_for_tree_label_with_parent(page, labels, parent_content=parent_content, timeout_ms=3000)
            if not branch:
                return False
        if not expand_tree_branch(page, labels, parent_content=parent_content):
            return expand_tree_until_label(page, path_segments[-1], root_labels=root_labels, timeout_ms=15000)
        parent_content = tree_node_content_from_label(branch)
        page.wait_for_timeout(600)

    final_leaf = wait_for_tree_label_with_parent(page, path_segments[-1], parent_content=parent_content, timeout_ms=5000)
    if not final_leaf:
        if not expand_tree_until_label(page, path_segments[-1], root_labels=root_labels, timeout_ms=15000):
            return False
        final_leaf = wait_for_tree_label_with_parent(page, path_segments[-1], parent_content=parent_content, timeout_ms=3000)
        if not final_leaf:
            return False
    return click_tree_leaf(page, path_segments[-1], force=True, parent_content=parent_content)


def open_route_from_tree(page, data):
    route_key = data.get("vec_route")
    route_config = ROUTE_LABELS.get(route_key)
    if not route_config:
        return page

    if not merit_listing_ready(page):
        raise RuntimeError("The merit listing page is not visible yet, so the left-side merit route cannot be selected.")

    if route_is_selected(page, route_config["selected"]):
        return bring_page_to_front(page)

    tree_segments = route_tree_segments(route_key, route_config)
    if not open_route_tree_path(page, tree_segments):
        target_label = tree_segments[-1][0] if tree_segments else route_config["tree_child"][0]
        raise RuntimeError(
            f"Could not open the merit type '{target_label}' "
            "while trying to open the next merit route."
        )

    page.wait_for_timeout(1200)
    listing_page = choose_start_page(page.context) or page
    if merit_listing_ready(listing_page) and wait_for_route_selection(
        listing_page,
        route_config["selected"],
        timeout_ms=10000,
    ):
        return bring_page_to_front(listing_page)

    target_label = tree_segments[-1][0] if tree_segments else route_config["tree_child"][0]
    current_label = current_route_label(listing_page) or current_route_label(page)
    raise RuntimeError(
        f"The left-side merit route '{target_label}' was opened, but the page did not confirm that selection. "
        f"Current selected route: '{current_label or 'unknown route'}'."
    )


def ensure_expected_route(page, data):
    route_key = data.get("vec_route")
    route_config = ROUTE_LABELS.get(route_key)
    if not route_config:
        return page

    if page_has_wizard_shell(page):
        return page

    current_label = current_route_label(page)
    if current_label and text_matches(current_label, route_config["selected"]):
        return page

    tree_segments = route_tree_segments(route_key, route_config)
    target_label = tree_segments[-1][0] if tree_segments else route_config["tree_child"][0]

    print(
        f"Current merit route '{current_label or 'unknown route'}' does not match expected vec_route "
        f"'{data.get('vec_route')}'. Trying to switch to '{target_label}'."
    )

    if open_route_tree_path(page, tree_segments):
        page = choose_start_page(page.context) or page
        if wait_for_route_selection(page, route_config["selected"], timeout_ms=12000):
            return page

        current_label = current_route_label(page)
    else:
        save_debug(page, "route_switch_missing_leaf")
        raise RuntimeError(
            f"Could not click the merit type '{target_label}' in the navigation tree. "
            f"The current route is '{current_label or 'unknown route'}'."
        )

    save_debug(page, "route_switch_failed")
    raise RuntimeError(
        f"Wrong merit type opened. Expected route '{target_label}' from classification_result, "
        f"but the current route is '{current_label or 'unknown route'}'."
    )


def attach_to_wizard_page(context, timeout_ms=30000):
    import time

    end = time.time() + (timeout_ms / 1000)

    while time.time() < end:
        pages = open_context_pages(context)
        for candidate in reversed(pages):
            for scope in iter_scopes(candidate):
                if scope_contains_markers(scope, WIZARD_STEP_MARKERS):
                    return bring_page_to_front(candidate)

        time.sleep(0.5)

    return choose_start_page(context)


def require_visible_step(context, step_name, step_markers, active_titles=None, field_labels=None):
    while True:
        page = attach_to_wizard_page(context, timeout_ms=10000)
        if page and wait_for_step(
            page,
            step_markers,
            active_titles=active_titles,
            field_labels=field_labels,
            timeout_ms=3000,
        ):
            return page

        print()
        print(f"Could not detect {step_name} yet.")
        pages = open_context_pages(context)
        if not pages:
            print("No open pages are available in the browser context right now.")
            print("If you closed the launched browser tab/window, reopen the VEC wizard in that same session or rerun the script.")
        elif len(pages) == 1 and safe_page_url(pages[0]) == PUBLIC_VEC_URL:
            print("The script can still only see the public VEC landing page.")
            print("Open the actual wizard inside the same browser window launched by this script, then retry.")
        else:
            print("Keep the VEC wizard visible in the same browser window, then retry.")
        save_context_debug(context, f"missing_{step_name}")
        input(f"When {step_name} is visible in that browser window, press Enter to retry...")


def fill_using_dom(scope, labels, value):
    script = """
([labels, value]) => {
  const normalize = (text) =>
    (text || "")
      .normalize("NFD")
      .replace(/[\\u0300-\\u036f]/g, "")
      .replace(/\\s+/g, " ")
      .trim()
      .toLowerCase();

  const wanted = labels.map(normalize).filter(Boolean);
  if (!wanted.length) {
    return false;
  }

  const isVisible = (element) => {
    if (!element) {
      return false;
    }
    const style = window.getComputedStyle(element);
    if (style.visibility === "hidden" || style.display === "none") {
      return false;
    }
    const rect = element.getBoundingClientRect();
    return rect.width > 0 && rect.height > 0;
  };

  const isField = (element) =>
    !!element &&
    ["INPUT", "TEXTAREA"].includes(element.tagName) &&
    element.type !== "hidden" &&
    !element.disabled;

  const setValue = (field) => {
    field.focus();
    field.value = "";
    field.dispatchEvent(new Event("input", { bubbles: true }));
    field.value = value;
    field.dispatchEvent(new Event("input", { bubbles: true }));
    field.dispatchEvent(new Event("change", { bubbles: true }));
  };

  const findVisibleField = (root) => {
    if (!root) {
      return null;
    }
    if (isField(root) && isVisible(root)) {
      return root;
    }
    return [...root.querySelectorAll("input:not([type='hidden']), textarea")]
      .find((field) => isField(field) && isVisible(field));
  };

  const textNodes = [...document.querySelectorAll("label, span, div, td, th, p, strong, legend")];
  for (const node of textNodes) {
    const text = normalize(node.textContent);
    if (!text || !wanted.some((label) => text.includes(label))) {
      continue;
    }

    const htmlFor = node.getAttribute && node.getAttribute("for");
    if (htmlFor) {
      const linked = document.getElementById(htmlFor);
      if (findVisibleField(linked)) {
        setValue(linked);
        return true;
      }
    }

    let container = node;
    for (let depth = 0; depth < 4 && container; depth += 1, container = container.parentElement) {
      const field = findVisibleField(container);
      if (field) {
        setValue(field);
        return true;
      }
    }

    let sibling = node.nextElementSibling;
    while (sibling) {
      const field = findVisibleField(sibling);
      if (field) {
        setValue(field);
        return true;
      }
      sibling = sibling.nextElementSibling;
    }
  }

  return false;
}
"""

    try:
        return bool(scope.evaluate(script, [labels, value]))
    except Exception:
        return False


def try_fill_locator(locator, value):
    try:
        locator.scroll_into_view_if_needed()
    except Exception:
        pass

    try:
        locator.fill("")
        locator.fill(value)
        return True
    except Exception:
        pass

    try:
        locator.click()
        locator.press("Control+A")
        locator.type(value, delay=20)
        return True
    except Exception:
        return False


def click_if_found(page, texts, timeout_ms=5000):
    for text in expand_ui_text_variants(texts):
        for scope in iter_scopes(page):
            locator = first_visible(
                [
                    scope.get_by_role("button", name=text, exact=False),
                    scope.get_by_role("link", name=text, exact=False),
                    scope.get_by_text(text, exact=False),
                ]
            )
            if not locator:
                continue

            try:
                locator.scroll_into_view_if_needed(timeout=timeout_ms)
            except Exception:
                pass

            try:
                locator.click(timeout=timeout_ms)
                page.wait_for_timeout(1500)
                return True
            except Exception:
                continue

    return False


def click_if_found_with_retry(page, texts, timeout_ms=7000, click_timeout_ms=2500):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        if click_if_found(page, texts, timeout_ms=click_timeout_ms):
            return True
        if not safe_wait_for_timeout(page, 300):
            return False

    return click_if_found(page, texts, timeout_ms=click_timeout_ms)


def fill_text_field(page, labels, value, required=False):
    if value in (None, ""):
        if required:
            raise RuntimeError(f"Required field missing for labels {labels}")
        return False

    value = str(value)

    for scope in iter_scopes(page):
        locator = first_visible(field_locators(scope, labels))
        if locator and try_fill_locator(locator, value):
            page.wait_for_timeout(400)
            return True

    for scope in iter_scopes(page):
        if fill_using_dom(scope, labels, value):
            page.wait_for_timeout(400)
            return True

    if required:
        debug_name = re.sub(r"[^a-zA-Z0-9_-]+", "_", "_".join(labels)).strip("_") or "missing_field"
        save_debug(page, f"missing_{debug_name}")
        raise RuntimeError(f"Could not find input for labels {labels} on the current wizard step.")
    return False


def input_value_matches(current_value, expected_value):
    normalized_current = normalize_text(current_value).replace(",", ".")
    normalized_expected = normalize_text(expected_value).replace(",", ".")

    if normalized_current == normalized_expected:
        return True

    try:
        return abs(float(normalized_current) - float(normalized_expected)) < 0.0001
    except Exception:
        return False


def current_text_field_value(page, labels):
    for scope in iter_scopes(page):
        locator = first_existing(field_locators(scope, labels))
        if not locator:
            continue

        try:
            return locator.input_value().strip()
        except Exception:
            pass

        try:
            return locator.inner_text().strip()
        except Exception:
            continue

    return ""


def wait_for_text_field_value(page, labels, expected_value, timeout_ms=5000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        current_value = current_text_field_value(page, labels)
        if input_value_matches(current_value, expected_value):
            return True
        if not safe_wait_for_timeout(page, 250):
            return False

    current_value = current_text_field_value(page, labels)
    return input_value_matches(current_value, expected_value)


def fill_verified_text_field(page, labels, value, required=False):
    if value in (None, ""):
        if required:
            raise RuntimeError(f"Required field missing for labels {labels}")
        return False

    value_text = str(value).strip()
    if not value_text:
        if required:
            raise RuntimeError(f"Required field missing for labels {labels}")
        return False

    if wait_for_text_field_value(page, labels, value_text, timeout_ms=800):
        return True

    if not fill_text_field(page, labels, value_text, required=required):
        return False

    if wait_for_text_field_value(page, labels, value_text, timeout_ms=5000):
        return True

    if required:
        current_value = current_text_field_value(page, labels) or "(blank)"
        raise RuntimeError(
            f"Could not lock the field {labels} to {value_text!r}. Current value is {current_value!r}."
        )

    return False


def select_from_autocomplete_if_present(page, preferred_text=None):
    for scope in iter_scopes(page):
        options = scope.locator(
            ".ui-autocomplete-item:visible, li[role='option']:visible, tr[role='row']:visible"
        )
        try:
            count = options.count()
        except Exception:
            count = 0

        if count == 0:
            continue

        if preferred_text:
            try:
                option = first_visible(
                    [
                        scope.get_by_text(preferred_text, exact=False),
                    ]
                )
                if option:
                    option.click(timeout=3000)
                    page.wait_for_timeout(800)
                    return True
            except Exception:
                pass

        try:
            options.first.click(timeout=3000)
            page.wait_for_timeout(800)
            return True
        except Exception:
            continue

    return False


def click_preferred_result(page, preferred_texts):
    cleaned_texts = [text for text in preferred_texts if text]

    for text in cleaned_texts:
        for scope in iter_scopes(page):
            candidates = [
                scope.get_by_role("link", name=text, exact=False),
                scope.get_by_role("button", name=text, exact=False),
                scope.get_by_text(text, exact=False),
                scope.locator(
                    f"xpath=(//tr[.//*[contains(normalize-space(.), \"{text}\")]])[1]"
                ),
                scope.locator(
                    f"xpath=(//*[contains(@class, 'ui-datatable') or contains(@class, 'ui-selectonelistbox') or contains(@class, 'ui-selectoneradio')]//*[contains(normalize-space(.), \"{text}\")])[1]"
                ),
            ]
            locator = first_visible(candidates)
            if not locator:
                continue

            try:
                locator.scroll_into_view_if_needed()
            except Exception:
                pass

            try:
                locator.click(timeout=5000)
                page.wait_for_timeout(1000)
                return True
            except Exception:
                continue

    for scope in iter_scopes(page):
        fallback = first_visible(
            [
                scope.locator("table tr:visible"),
                scope.locator("[role='row']:visible"),
                scope.locator(".ui-datatable-data tr:visible"),
                scope.locator("li[role='option']:visible"),
            ]
        )
        if not fallback:
            continue

        try:
            fallback.scroll_into_view_if_needed()
        except Exception:
            pass

        try:
            fallback.click(timeout=5000)
            page.wait_for_timeout(1000)
            return True
        except Exception:
            continue

    return False


def upload_pdf(page, pdf_path):
    for scope in iter_scopes(page):
        file_input = scope.locator("input[type='file']")
        try:
            if file_input.count() > 0:
                file_input.first.set_input_files(pdf_path)
                page.wait_for_timeout(1200)
                return True
        except Exception:
            continue

    click_if_found(
        page,
        [
            "Añadir documento",
            "Anadir documento",
            "Subir archivo",
            "Subir fichero",
            "Adjuntar",
            "Attach",
            "Upload",
        ],
        timeout_ms=4000,
    )

    for scope in iter_scopes(page):
        file_input = scope.locator("input[type='file']")
        try:
            if file_input.count() > 0:
                file_input.first.set_input_files(pdf_path)
                page.wait_for_timeout(1200)
                return True
        except Exception:
            continue

    return False


def accept_modal_if_present(page, timeout_ms=5000):
    direct_locators = []
    for scope in iter_scopes(page):
        direct_locators.extend(
            [
                scope.locator("#form\\:guardarSigMeritoPanelAceptar"),
                scope.locator("button[id$='guardarSigMeritoPanelAceptar']"),
                scope.locator("input[id$='guardarSigMeritoPanelAceptar']"),
            ]
        )

    locator = first_visible(direct_locators)
    if locator:
        try:
            locator.click(timeout=timeout_ms)
            page.wait_for_timeout(1500)
            return True
        except Exception:
            pass

    return click_if_found(page, ["Aceptar", "Accept", "Confirmar", "Confirm"], timeout_ms=timeout_ms)


def drain_accept_modals(page, timeout_ms=5000):
    import time

    end = time.time() + (timeout_ms / 1000)
    clicked_any = False

    while time.time() < end:
        clicked = accept_modal_if_present(page, timeout_ms=1200)
        if clicked:
            clicked_any = True
            continue

        if clicked_any:
            break

        if not safe_wait_for_timeout(page, 250):
            break

    return clicked_any


def submit_and_wait_for_step(page, target_step_name, active_titles, submitter, timeout_ms=12000, max_attempts=3):
    wait_per_attempt_ms = max(2500, int(timeout_ms / max(1, max_attempts)))

    for _ in range(max_attempts):
        drain_accept_modals(page, timeout_ms=1500)

        if submitter():
            if wait_for_active_step(page, active_titles, timeout_ms=wait_per_attempt_ms):
                return True

        modal_clicked = drain_accept_modals(page, timeout_ms=3500)
        if modal_clicked and wait_for_active_step(page, active_titles, timeout_ms=wait_per_attempt_ms):
            return True

        if wait_for_active_step(page, active_titles, timeout_ms=1500):
            return True

        safe_wait_for_timeout(page, 700)

    ensure_transition_to_step(page, target_step_name, active_titles, timeout_ms=timeout_ms)
    fill_step4_input_by_labels(
        page,
        ["CÃ³digo Curso", "Codigo curso", "Course code"],
        data.get("course_code"),
    )

    return True
    fill_step4_input_by_labels(
        page,
        ["CÃ³digo Curso", "Codigo curso", "Course code"],
        data.get("course_code"),
    )

    fill_step4_input_by_labels(
        page,
        ["CÃ³digo Curso", "Codigo curso", "Course code"],
        data.get("course_code"),
    )

    return True


def click_next(page, required=True):
    direct_locators = []
    for scope in iter_scopes(page):
        direct_locators.extend(
            [
                scope.locator("#form\\:guardarSigButton"),
                scope.locator("button[id$='guardarSigButton']"),
                scope.locator("#form\\:btnSiguiente"),
                scope.locator("button[id$='btnSiguiente']"),
                scope.locator("input[id$='btnSiguiente']"),
            ]
        )

    locator = first_visible(direct_locators)
    ok = False
    if locator:
        try:
            locator.scroll_into_view_if_needed(timeout=7000)
        except Exception:
            pass

        try:
            locator.click(timeout=7000)
            page.wait_for_timeout(1500)
            accept_modal_if_present(page)
            ok = True
        except Exception:
            ok = False

    if not ok:
        ok = click_if_found(page, ["Guardar y Siguiente", "Siguiente", "Next"], timeout_ms=7000)
        if ok:
            accept_modal_if_present(page)
    if not ok and required:
        raise RuntimeError("Could not click Next / Siguiente")
    return ok


def submit_step4(page):
    direct_locators = []
    for scope in iter_scopes(page):
        direct_locators.extend(
            [
                scope.locator("#form\\:guardarSigButton"),
                scope.locator("button[id$='guardarSigButton']"),
                scope.locator("#form\\:btnSiguiente"),
                scope.locator("button[id$='btnSiguiente']"),
                scope.locator("input[id$='btnSiguiente']"),
            ]
        )

    locator = first_visible(direct_locators)
    if locator:
        try:
            locator.scroll_into_view_if_needed(timeout=7000)
        except Exception:
            pass

        try:
            locator.click(timeout=7000)
            page.wait_for_timeout(1500)
            accept_modal_if_present(page)
            return True
        except Exception:
            pass

    if click_if_found(page, ["Guardar y Siguiente", "Save and Next", "Continuar", "Continue", "Siguiente", "Next"], timeout_ms=7000):
        accept_modal_if_present(page)
        return True

    return click_next(page, required=False)


def step_has_visible_controls(page, field_labels):
    if not field_labels:
        return True

    for labels in field_labels:
        for scope in iter_scopes(page):
            if first_visible(field_locators(scope, labels)):
                return True
            if first_visible(select_option_locators(scope, labels)):
                return True

    return False


def page_has_wizard_shell(page):
    if current_active_step_title(page):
        return True

    for scope in iter_scopes(page):
        candidates = [
            scope.locator(".ui-steps"),
            scope.locator(".ui-steps-item"),
            scope.locator("#form\\:guardarSigButton"),
            scope.locator("button[id$='guardarSigButton']"),
            scope.locator("#form\\:btnSiguiente"),
            scope.locator("button[id$='btnSiguiente']"),
            scope.locator("input[id$='btnSiguiente']"),
        ]
        if first_existing(candidates):
            return True

        if scope_contains_markers(scope, ["BÃºsqueda de MÃ©ritos", "Busqueda de Meritos"]):
            return False

    return False


def page_matches_step(page, step_markers, active_titles=None, field_labels=None):
    active_title = current_active_step_title(page)
    if active_titles:
        if active_title and text_matches(active_title, active_titles):
            return True
    elif active_title:
        return True

    marker_found = False
    for scope in iter_scopes(page):
        if scope_contains_markers(scope, step_markers):
            marker_found = True
            break

    if not marker_found:
        return False
    if active_titles and active_title:
        return False
    if active_titles and not page_has_wizard_shell(page):
        return False
    if field_labels and not step_has_visible_controls(page, field_labels):
        return False
    return True


def wait_for_step(page, step_markers, active_titles=None, field_labels=None, timeout_ms=15000):
    import time

    end = time.time() + (timeout_ms / 1000)

    while time.time() < end:
        if page_matches_step(page, step_markers, active_titles=active_titles, field_labels=field_labels):
            return True
        if not safe_wait_for_timeout(page, 500):
            return False

    return page_matches_step(page, step_markers, active_titles=active_titles, field_labels=field_labels)


def ensure_transition_to_step(page, step_name, active_titles, timeout_ms=12000):
    if wait_for_active_step(page, active_titles, timeout_ms=timeout_ms):
        return

    current_step = current_active_step_title(page) or "unknown step"
    error_text = "; ".join(visible_error_messages(page)) or "no visible validation message"
    debug_name = re.sub(r"[^a-zA-Z0-9_-]+", "_", step_name).strip("_") or "transition_error"
    save_debug(page, f"transition_failed_{debug_name}")
    raise RuntimeError(
        f"Did not reach {step_name}. Current step is '{current_step}'. Visible errors: {error_text}"
    )


def get_center_selection_value(page):
    for scope in iter_scopes(page):
        locator = first_existing(
            [
                scope.locator("#form\\:resultTbl_selection"),
                scope.locator("input[id$='resultTbl_selection']"),
            ]
        )
        if not locator:
            continue

        try:
            value = locator.input_value().strip()
        except Exception:
            value = (locator.get_attribute("value") or "").strip()

        if value:
            return value

    return ""


def center_is_selected(page):
    if get_center_selection_value(page):
        return True

    for scope in iter_scopes(page):
        selected_row = first_visible(
            [
                scope.locator("#form\\:resultTbl_data tr[aria-selected='true']"),
                scope.locator("tbody[id$='resultTbl_data'] tr[aria-selected='true']"),
                scope.locator("#form\\:resultTbl_data tr.ui-state-highlight"),
                scope.locator("tbody[id$='resultTbl_data'] tr.ui-state-highlight"),
            ]
        )
        if selected_row:
            return True

    return False


def wait_for_center_selected(page, timeout_ms=8000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        if center_is_selected(page):
            return True
        page.wait_for_timeout(300)

    return center_is_selected(page)


def dismiss_center_autocomplete(page):
    for scope in iter_scopes(page):
        locator = first_visible(field_locators(scope, ["Centro", "Center"]))
        if not locator:
            continue

        for key in ("Escape", "Tab"):
            try:
                locator.press(key)
                page.wait_for_timeout(150)
            except Exception:
                continue


def center_source_texts(data):
    def compact_text(value):
        return re.sub(r"\s+", " ", str(value or "")).strip(" :;,.\"'")

    sources = [
        data.get("center_search_text"),
        data.get("institution_name"),
        data.get("organizing_entity"),
        data.get("center_name"),
        data.get("center_result_text"),
    ]

    field_evidence = data.get("field_evidence")
    if isinstance(field_evidence, dict):
        for key in (
            "institution_name",
            "organizing_entity",
            "center_name",
            "center_search_text",
            "center_result_text",
        ):
            sources.append(field_evidence.get(key))

    result = []
    seen = set()
    for value in sources:
        compact = compact_text(value)
        normalized = normalize_text(compact)
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        result.append(compact)

    return result


def extract_center_terms_from_text(value):
    def compact_text(text):
        return re.sub(r"\s+", " ", str(text or "")).strip(" :;,.\"'")

    text = compact_text(value)
    if not text:
        return []

    repaired_text = repair_mojibake_text(text)
    candidates = []
    generic_tokens = {
        "actividad",
        "docente",
        "formacion",
        "continuada",
        "educativa",
        "centro",
        "centros",
        "salud",
        "sanitaria",
        "sanitario",
        "hospital",
        "hospitalaria",
        "universidad",
        "universitario",
        "universitaria",
        "asociacion",
        "federacion",
        "colegio",
        "oficial",
        "servicio",
        "andaluz",
        "ensenanza",
        "documento",
        "firmado",
        "comision",
    }
    patterns = [
        r"organizad[oa]\s+por\s+(.+?)(?:,?\s+celebrad[oa]|,?\s+realizad[oa]|,|\.)",
        r"documento firmado por\s+(.+?)(?:\s+-\s+[A-Z0-9]{6,}|,|\.)",
        r"firmad[oa]\s+digitalmente\s+por\s+(.+?)(?:\s+\(|,|\.)",
        r"\bo=([A-ZÁÉÍÓÚÑ0-9 .,&\-]{6,}?)(?:,|\)|$)",
    ]

    for pattern in patterns:
        for match in re.finditer(pattern, repaired_text, flags=re.IGNORECASE | re.DOTALL):
            candidate = compact_text(match.group(1) if match.lastindex else match.group(0))
            candidate = re.sub(r"^(?:la|el)\s+", "", candidate, flags=re.IGNORECASE).strip()
            candidate = re.sub(r"^(?:por|organizad[oa]\s+por)\s+", "", candidate, flags=re.IGNORECASE).strip()
            candidate = re.sub(r"\s+-\s+[A-Z0-9]{6,}$", "", candidate).strip()
            if candidate:
                candidates.append(candidate)

    cleaned_text = compact_text(
        re.sub(
            r",?\s+\b(?:S\.?\s*L(?:\.?\s*U\.)?|S\.?\s*A\.?|S\.?\s*C\.?\s*A\.?|S\.?\s*L\.?\s*L\.?)\b.*$",
            "",
            repaired_text,
            flags=re.IGNORECASE,
        )
    )
    if cleaned_text and normalize_text(cleaned_text) != normalize_text(text):
        candidates.append(cleaned_text)

    for part in re.split(r"[;,\n|]", repaired_text):
        part = compact_text(part)
        if part and len(part.split()) <= 8:
            candidates.append(part)

    for token in re.findall(r"\b[A-ZÁÉÍÓÚÑ][A-Za-zÁÉÍÓÚÑáéíóúñ]{3,}\b", repaired_text):
        if normalize_text(token) not in generic_tokens:
            candidates.append(token)

    if len(text.split()) <= 6 and len(text) <= 80:
        candidates.append(text)

    filtered_candidates = []
    for candidate in candidates:
        candidate = re.sub(r"^(?:por|organizad[oa]\s+por)\s+", "", str(candidate or ""), flags=re.IGNORECASE).strip()
        if not candidate:
            continue
        normalized_letters = re.sub(r"[^a-z0-9]+", "", normalize_text(candidate))
        if normalized_letters in {"sl", "slu", "sa", "sca", "sll"}:
            continue
        filtered_candidates.append(candidate)

    return dedupe_preserve_order(filtered_candidates)


def best_center_description_text(data):
    direct_values = [
        data.get("institution_name"),
        data.get("organizing_entity"),
        data.get("center_name"),
        data.get("center_search_text"),
    ]
    for value in direct_values:
        cleaned = re.sub(r"\s+", " ", str(value or "")).strip()
        if cleaned:
            return cleaned

    for source in center_source_texts(data):
        extracted_terms = extract_center_terms_from_text(source)
        if extracted_terms:
            return extracted_terms[0]

    return ""


def build_center_result_preferences(data):
    preferences = []
    for source in center_source_texts(data):
        preferences.extend(extract_center_terms_from_text(source))

    institution = normalize_text(" ".join(center_source_texts(data)))

    if any(token in institution for token in ("asociacion", "association", "federacion")):
        preferences.extend(["Asociación", "Asociacion", "Federación", "Federacion"])

    if any(token in institution for token in ("academia", "formacion", "ensenanza", "docencia")):
        preferences.extend(["ACADEMIAS", "ENSEÑANZA", "ENSENANZA"])

    if any(token in institution for token in ("universidad", "universitario")):
        preferences.extend(
            [
                "UNIVERSITARIA",
                "UNIVERSITARIO",
                "UNIVERSIDAD",
                "UNIVERSIDADES",
                "CENTRO UNIVERSITARIO",
                "CENTROS UNIVERSITARIOS",
            ]
        )

    if "servicio andaluz de salud" in institution or re.search(r"\bsas\b", institution):
        preferences.extend(["SERVICIO ANDALUZ DE SALUD", "SAS"])

    if any(token in institution for token in ("hospital", "clinica", "sanitari", "salud")):
        preferences.extend(["SANITARIAS", "SANITARIA", "HOSPITALARIAS", "HOSPITALARIA"])

    return dedupe_preserve_order(preferences)


def build_center_search_terms(data):
    terms = []
    for source in center_source_texts(data):
        terms.extend(extract_center_terms_from_text(source))

    center_result_text = re.sub(r"\s+", " ", str(data.get("center_result_text") or "")).strip(" :;,.\"'")
    if center_result_text and "-" in center_result_text:
        terms.append(center_result_text.split("-", 1)[1].strip())

    institution_name = best_center_description_text(data)
    normalized_institution = normalize_text(institution_name)

    if any(token in normalized_institution for token in ("universidad", "universitario")):
        terms.extend(["Universidad", "Universidades", "Centro Universitario", "Centros Universitarios"])
        match = re.search(
            r"\buniversidad(?:\s+de|\s+del|\s+internacional de)?\s+(.+)",
            institution_name,
            flags=re.IGNORECASE,
        )
        if match:
            terms.append(compact_text(match.group(1)))

    if "servicio andaluz de salud" in normalized_institution or re.search(r"\bsas\b", normalized_institution):
        terms.extend(["SAS", "Servicio Andaluz de Salud"])

    if any(token in normalized_institution for token in ("colegio oficial", "colegio ")):
        terms.extend(["Colegio Oficial", "Colegio"])

    if any(token in normalized_institution for token in ("asociacion", "association", "federacion")):
        terms.extend(["Asociación", "Federación"])

    if any(token in normalized_institution for token in ("hospital", "clinica", "sanitari", "salud")):
        terms.extend(["Hospital", "Clínica", "Centro Sanitario"])

    return dedupe_preserve_order(terms)


def _unused_build_center_result_preferences_override_2(data):
    preferences = []
    for source in center_source_texts(data):
        preferences.extend(extract_center_terms_from_text(source))

    institution = normalize_text(" ".join(center_source_texts(data)))
    center_result_text = normalize_text(data.get("center_result_text"))
    association_center_requested = center_result_text.startswith("02-") or any(
        token in center_result_text for token in ("asociacion", "association", "federacion", "federation")
    )

    if association_center_requested or any(token in institution for token in ("asociacion", "association", "federacion")):
        preferences.extend(["AsociaciÃ³n", "Asociacion", "FederaciÃ³n", "Federacion"])

    if any(token in institution for token in ("academia", "formacion", "ensenanza", "docencia")):
        preferences.extend(["ACADEMIAS", "ENSEÃ‘ANZA", "ENSENANZA"])

    if any(token in institution for token in ("universidad", "universitario")):
        preferences.extend(
            [
                "UNIVERSITARIA",
                "UNIVERSITARIO",
                "UNIVERSIDAD",
                "UNIVERSIDADES",
                "CENTRO UNIVERSITARIO",
                "CENTROS UNIVERSITARIOS",
            ]
        )

    if not association_center_requested and ("servicio andaluz de salud" in institution or re.search(r"\bsas\b", institution)):
        preferences.extend(["SERVICIO ANDALUZ DE SALUD", "SAS"])

    if not association_center_requested and any(token in institution for token in ("hospital", "clinica", "sanitari", "salud")):
        preferences.extend(["SANITARIAS", "SANITARIA", "HOSPITALARIAS", "HOSPITALARIA"])

    return dedupe_preserve_order(preferences)


def _unused_build_center_search_terms_override_2(data):
    terms = []
    for source in center_source_texts(data):
        terms.extend(extract_center_terms_from_text(source))

    center_result_text = re.sub(r"\s+", " ", str(data.get("center_result_text") or "")).strip(" :;,.\"'")
    if center_result_text and "-" in center_result_text:
        terms.append(center_result_text.split("-", 1)[1].strip())

    institution_name = best_center_description_text(data)
    normalized_institution = normalize_text(institution_name)
    normalized_center_result = normalize_text(data.get("center_result_text"))
    association_center_requested = normalized_center_result.startswith("02-") or any(
        token in normalized_center_result for token in ("asociacion", "association", "federacion", "federation")
    )

    if any(token in normalized_institution for token in ("universidad", "universitario")):
        terms.extend(["Universidad", "Universidades", "Centro Universitario", "Centros Universitarios"])
        match = re.search(
            r"\buniversidad(?:\s+de|\s+del|\s+internacional de)?\s+(.+)",
            institution_name,
            flags=re.IGNORECASE,
        )
        if match:
            terms.append(compact_text(match.group(1)))

    if not association_center_requested and (
        "servicio andaluz de salud" in normalized_institution or re.search(r"\bsas\b", normalized_institution)
    ):
        terms.extend(["SAS", "Servicio Andaluz de Salud"])

    if any(token in normalized_institution for token in ("colegio oficial", "colegio ")):
        terms.extend(["Colegio Oficial", "Colegio"])

    if association_center_requested or any(
        token in normalized_institution for token in ("asociacion", "association", "federacion", "federation")
    ):
        terms.extend(["AsociaciÃ³n", "FederaciÃ³n"])

    if not association_center_requested and any(token in normalized_institution for token in ("hospital", "clinica", "sanitari", "salud")):
        terms.extend(["Hospital", "ClÃ­nica", "Centro Sanitario"])

    return dedupe_preserve_order(terms)


def _unused_build_center_result_preferences_override_3(data):
    preferences = []
    source_texts = center_source_texts(data)
    institution = normalize_text(" ".join(source_texts))
    center_result_text = normalize_text(data.get("center_result_text"))
    association_center_requested = center_result_text.startswith("02-") or any(
        token in center_result_text for token in ("asociacion", "association", "federacion", "federation")
    )

    if association_center_requested:
        preferences.extend(["AsociaciÃ³n", "Asociacion", "FederaciÃ³n", "Federacion"])

    for source in source_texts:
        normalized_source = normalize_text(source)
        if association_center_requested and (
            "servicio andaluz de salud" in normalized_source or re.search(r"\bsas\b", normalized_source)
        ):
            continue
        preferences.extend(extract_center_terms_from_text(source))

    if any(token in institution for token in ("asociacion", "association", "federacion")):
        preferences.extend(["AsociaciÃ³n", "Asociacion", "FederaciÃ³n", "Federacion"])

    if any(token in institution for token in ("academia", "formacion", "ensenanza", "docencia")):
        preferences.extend(["ACADEMIAS", "ENSEÃ‘ANZA", "ENSENANZA"])

    if any(token in institution for token in ("universidad", "universitario")):
        preferences.extend(
            [
                "UNIVERSITARIA",
                "UNIVERSITARIO",
                "UNIVERSIDAD",
                "UNIVERSIDADES",
                "CENTRO UNIVERSITARIO",
                "CENTROS UNIVERSITARIOS",
            ]
        )

    if not association_center_requested and ("servicio andaluz de salud" in institution or re.search(r"\bsas\b", institution)):
        preferences.extend(["SERVICIO ANDALUZ DE SALUD", "SAS"])

    if not association_center_requested and any(token in institution for token in ("hospital", "clinica", "sanitari", "salud")):
        preferences.extend(["SANITARIAS", "SANITARIA", "HOSPITALARIAS", "HOSPITALARIA"])

    return dedupe_preserve_order(preferences)


def _unused_build_center_search_terms_override_3(data):
    terms = []
    center_result_text = re.sub(r"\s+", " ", str(data.get("center_result_text") or "")).strip(" :;,.\"'")
    if center_result_text and "-" in center_result_text:
        terms.append(center_result_text.split("-", 1)[1].strip())

    source_texts = center_source_texts(data)
    institution_name = best_center_description_text(data)
    normalized_institution = normalize_text(institution_name)
    normalized_center_result = normalize_text(data.get("center_result_text"))
    association_center_requested = normalized_center_result.startswith("02-") or any(
        token in normalized_center_result for token in ("asociacion", "association", "federacion", "federation")
    )

    for source in source_texts:
        normalized_source = normalize_text(source)
        if association_center_requested and (
            "servicio andaluz de salud" in normalized_source or re.search(r"\bsas\b", normalized_source)
        ):
            continue
        terms.extend(extract_center_terms_from_text(source))

    if any(token in normalized_institution for token in ("universidad", "universitario")):
        terms.extend(["Universidad", "Universidades", "Centro Universitario", "Centros Universitarios"])
        match = re.search(
            r"\buniversidad(?:\s+de|\s+del|\s+internacional de)?\s+(.+)",
            institution_name,
            flags=re.IGNORECASE,
        )
        if match:
            terms.append(compact_text(match.group(1)))

    if not association_center_requested and (
        "servicio andaluz de salud" in normalized_institution or re.search(r"\bsas\b", normalized_institution)
    ):
        terms.extend(["SAS", "Servicio Andaluz de Salud"])

    if any(token in normalized_institution for token in ("colegio oficial", "colegio ")):
        terms.extend(["Colegio Oficial", "Colegio"])

    if association_center_requested or any(
        token in normalized_institution for token in ("asociacion", "association", "federacion", "federation")
    ):
        terms.extend(["AsociaciÃ³n", "FederaciÃ³n"])

    if not association_center_requested and any(token in normalized_institution for token in ("hospital", "clinica", "sanitari", "salud")):
        terms.extend(["Hospital", "ClÃ­nica", "Centro Sanitario"])

    return dedupe_preserve_order(terms)


def center_result_rows(page):
    rows = []
    for scope in iter_scopes(page):
        try:
            locator = scope.locator(
                "#form\\:resultTbl_data tr.ui-datatable-selectable, tbody[id$='resultTbl_data'] tr.ui-datatable-selectable"
            )
            count = locator.count()
        except Exception:
            continue

        for index in range(count):
            row = locator.nth(index)
            try:
                if row.is_visible():
                    rows.append(row)
            except Exception:
                continue

    return rows


def click_center_result(page, data):
    preferences = build_center_result_preferences(data)
    rows = center_result_rows(page)
    if not rows:
        return False

    best_row = None
    best_score = 0
    for row in rows:
        try:
            row_text = row.inner_text().strip()
        except Exception:
            continue

        score = text_match_score(row_text, preferences)
        if score > best_score:
            best_row = row
            best_score = score

    if not best_row and len(rows) == 1:
        best_row = rows[0]
    if not best_row:
        return False

    click_targets = [best_row]
    try:
        first_cell = best_row.locator("td").first
        if first_cell.count() > 0:
            click_targets.append(first_cell)
    except Exception:
        pass

    for target in click_targets:
        try:
            target.scroll_into_view_if_needed()
        except Exception:
            pass

        for click_kwargs in ({}, {"force": True}):
            try:
                target.click(timeout=5000, **click_kwargs)
                page.wait_for_timeout(800)
            except Exception:
                try:
                    target.dispatch_event("click")
                    page.wait_for_timeout(800)
                except Exception:
                    continue

            if wait_for_center_selected(page, timeout_ms=4000):
                return True

    return False


def wait_for_center_results(page, timeout_ms=10000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        if center_result_rows(page):
            return True
        page.wait_for_timeout(300)

    return bool(center_result_rows(page))


def select_option_locators(scope, labels):
    locators = []
    for label in expand_ui_text_variants(labels):
        locators.extend(
            [
                scope.locator(
                    f"xpath=(//*[self::label or self::span or self::div or self::td or self::th][contains(normalize-space(.), \"{label}\")]/following::*[contains(@class, 'ui-selectonemenu')][1])[1]"
                ),
                scope.locator(
                    f"xpath=(//*[contains(normalize-space(.), \"{label}\")]/ancestor::*[self::div or self::td or self::th or self::tr][1]//*[contains(@class, 'ui-selectonemenu')][1])[1]"
                ),
                scope.locator(
                    f"xpath=(//*[self::label or self::span or self::div or self::td or self::th][contains(normalize-space(.), \"{label}\")]/following::*[self::select][1])[1]"
                ),
                scope.locator(
                    f"xpath=(//*[contains(normalize-space(.), \"{label}\")]/ancestor::*[self::div or self::td or self::th or self::tr][1]//*[self::select][1])[1]"
                ),
            ]
        )
    return locators


def select_option_candidates(value, field_kind=None):
    candidates = [str(value)] if value not in (None, "") else []
    normalized = normalize_text(value)

    if field_kind == "delivery_method":
        if normalized in ("online", "on line", "distancia", "a distancia", "teleformacion"):
            candidates.extend(["Online", "On line", "No Presencial", "A distancia", "Teleformación", "Teleformacion"])
        elif normalized in ("presencial",):
            candidates.extend(["Presencial"])
        elif normalized in ("mixta", "semipresencial"):
            candidates.extend(["Mixta", "Semipresencial"])

    return dedupe_preserve_order(candidates)


def try_select_hidden_select(select_locator, option_texts):
    try:
        options = select_locator.locator("option")
        count = options.count()
    except Exception:
        return False

    best_value = None
    best_score = 0
    for index in range(count):
        option = options.nth(index)
        try:
            option_text = option.inner_text().strip()
            option_value = (option.get_attribute("value") or "").strip()
        except Exception:
            continue

        if not option_value:
            continue

        score = text_match_score(option_text, option_texts)
        if score > best_score:
            best_score = score
            best_value = option_value

    if not best_value:
        return False

    try:
        select_locator.select_option(value=best_value)
        try:
            select_locator.evaluate(
                """(element) => {
                    element.dispatchEvent(new Event("input", { bubbles: true }));
                    element.dispatchEvent(new Event("change", { bubbles: true }));
                    element.dispatchEvent(new Event("blur", { bubbles: true }));
                }"""
            )
        except Exception:
            pass
        return True
    except Exception:
        pass

    try:
        select_locator.evaluate(
            """(element, value) => {
                element.value = value;
                element.dispatchEvent(new Event("input", { bubbles: true }));
                element.dispatchEvent(new Event("change", { bubbles: true }));
            }""",
            best_value,
        )
        return True
    except Exception:
        return False


def selected_option_text_from_menu(menu_locator):
    label_locator = first_existing(
        [
            menu_locator.locator("label.ui-selectonemenu-label"),
            menu_locator.locator(".ui-selectonemenu-label"),
        ]
    )
    if label_locator:
        try:
            label_text = label_locator.inner_text().strip()
        except Exception:
            label_text = ""

        if label_text:
            return label_text

    select_locator = first_existing(
        [
            menu_locator.locator("select"),
            menu_locator.locator("xpath=.//select[1]"),
            menu_locator,
        ]
    )
    if select_locator:
        try:
            selected_option = select_locator.locator("option:checked")
            if selected_option.count() > 0:
                return selected_option.first.inner_text().strip()
        except Exception:
            pass

    return ""


def option_text_matches_selection(selected_text, option_texts):
    if not selected_text or text_matches(selected_text, ["Selecciona", "Select"]):
        return False

    return text_match_score(selected_text, option_texts) > 0


def current_option_field_text(page, labels):
    for scope in iter_scopes(page):
        menu = first_visible(select_option_locators(scope, labels))
        if not menu:
            menu = first_existing(select_option_locators(scope, labels))
        if not menu:
            continue

        selected_text = selected_option_text_from_menu(menu)
        if selected_text:
            return selected_text

    return ""


def wait_for_option_field_value(page, labels, option_texts, timeout_ms=6000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        selected_text = current_option_field_text(page, labels)
        if option_text_matches_selection(selected_text, option_texts):
            return True
        if not safe_wait_for_timeout(page, 250):
            return False

    selected_text = current_option_field_text(page, labels)
    return option_text_matches_selection(selected_text, option_texts)


def try_select_menu_option(page, menu_locator, option_texts, labels=None):
    hidden_select = first_existing(
        [
            menu_locator.locator("select"),
            menu_locator.locator("xpath=.//select[1]"),
        ]
    )
    if hidden_select and try_select_hidden_select(hidden_select, option_texts):
        page.wait_for_timeout(700)
        if not labels or wait_for_option_field_value(page, labels, option_texts, timeout_ms=4000):
            return True

    try:
        menu_locator.scroll_into_view_if_needed()
    except Exception:
        pass

    click_target = first_visible(
        [
            menu_locator.locator(".ui-selectonemenu-trigger"),
            menu_locator.locator("label.ui-selectonemenu-label"),
            menu_locator,
        ]
    ) or menu_locator

    try:
        click_target.click(timeout=5000)
        page.wait_for_timeout(700)
    except Exception:
        return False

    best_item = None
    best_score = 0
    for scope in iter_scopes(page):
        try:
            items = scope.locator(
                ".ui-selectonemenu-panel:visible li.ui-selectonemenu-item, li.ui-selectonemenu-item"
            )
            count = items.count()
        except Exception:
            continue

        for index in range(count):
            item = items.nth(index)
            try:
                if not item.is_visible():
                    continue
                item_text = item.inner_text().strip()
            except Exception:
                continue

            if text_matches(item_text, ["Selecciona", "Select"]):
                continue

            score = text_match_score(item_text, option_texts)
            if score > best_score:
                best_item = item
                best_score = score

    if not best_item:
        return False

    try:
        best_item.click(timeout=5000)
        page.wait_for_timeout(700)
        if not labels:
            return True
        return wait_for_option_field_value(page, labels, option_texts, timeout_ms=5000)
    except Exception:
        return False


def select_option_field(page, labels, value, required=False, field_kind=None):
    option_texts = select_option_candidates(value, field_kind=field_kind)
    if not option_texts:
        if required:
            raise RuntimeError(f"Required selection missing for labels {labels}")
        return False

    if wait_for_option_field_value(page, labels, option_texts, timeout_ms=800):
        return True

    for scope in iter_scopes(page):
        menu = first_visible(select_option_locators(scope, labels))
        if menu and try_select_menu_option(page, menu, option_texts, labels=labels):
            return True

    if required:
        debug_name = re.sub(r"[^a-zA-Z0-9_-]+", "_", "_".join(labels)).strip("_") or "missing_select"
        save_debug(page, f"missing_{debug_name}")
        selected_text = current_option_field_text(page, labels) or "(blank)"
        raise RuntimeError(
            f"Could not select an option for labels {labels} on the current wizard step. "
            f"Current value is {selected_text!r}."
        )
    return False


def step4_row_locator(scope, row_index, suffix):
    return scope.locator(f"xpath=//*[@id='form:tblFields:{row_index}:idSelectOneMenu{suffix}']")


def step4_row_menu_exists(page, row_index):
    for scope in iter_scopes(page):
        if first_existing(
            [
                step4_row_locator(scope, row_index, ""),
                step4_row_locator(scope, row_index, "_input"),
            ]
        ):
            return True

    return False


def step4_row_records(page):
    records = []
    seen = set()

    for scope in iter_scopes(page):
        try:
            rows = scope.locator("xpath=//*[@id='form:tblFields_data']/tr")
            count = rows.count()
        except Exception:
            continue

        for index in range(count):
            row = rows.nth(index)
            row_index = ""
            try:
                row_index = (row.get_attribute("data-rk") or "").strip()
            except Exception:
                row_index = ""

            if not row_index:
                try:
                    row_id = (row.get_attribute("id") or "").strip()
                except Exception:
                    row_id = ""
                match = re.search(r"form:tblFields_node_(\d+)$", row_id)
                if match:
                    row_index = match.group(1)

            if not row_index or row_index in seen:
                continue

            label_text = ""
            try:
                labels = row.locator("label")
                if labels.count() > 0:
                    label_text = labels.first.inner_text().strip()
            except Exception:
                label_text = ""

            if not label_text:
                try:
                    label_cell = row.locator("td").first
                    if label_cell.count() > 0:
                        label_text = label_cell.inner_text().strip()
                except Exception:
                    label_text = ""

            if not label_text:
                continue

            seen.add(row_index)
            records.append(
                {
                    "row_index": int(row_index),
                    "label": label_text,
                }
            )

    records.sort(key=lambda item: item["row_index"])
    return records


def find_step4_row_index(page, labels):
    row_records = step4_row_records(page)
    if not row_records:
        return None

    label_candidates = expand_ui_text_variants(labels)
    best_row_index = None
    best_score = 0

    for record in row_records:
        score = text_match_score(record["label"], label_candidates)
        if score > best_score:
            best_score = score
            best_row_index = record["row_index"]

    if best_score <= 0:
        return None

    return best_row_index


def wait_for_step4_row_index(page, labels, timeout_ms=4000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        row_index = find_step4_row_index(page, labels)
        if row_index is not None:
            return row_index
        if not safe_wait_for_timeout(page, 250):
            return None

    return find_step4_row_index(page, labels)


def current_step4_row_text(page, row_index):
    for scope in iter_scopes(page):
        label_locator = first_existing(
            [
                step4_row_locator(scope, row_index, "_label"),
            ]
        )
        if not label_locator:
            continue

        try:
            label_text = label_locator.inner_text().strip()
        except Exception:
            label_text = ""

        if label_text:
            return label_text

    return ""


def wait_for_step4_row_value(page, row_index, option_texts, timeout_ms=6000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        selected_text = current_step4_row_text(page, row_index)
        if option_text_matches_selection(selected_text, option_texts):
            return True
        if not safe_wait_for_timeout(page, 250):
            return False

    selected_text = current_step4_row_text(page, row_index)
    return option_text_matches_selection(selected_text, option_texts)


def select_step4_row_option(page, row_index, value, required=False, field_kind=None):
    option_texts = select_option_candidates(value, field_kind=field_kind)
    if not option_texts:
        if required:
            raise RuntimeError(f"Required Step 4 row {row_index} selection is missing.")
        return False

    if wait_for_step4_row_value(page, row_index, option_texts, timeout_ms=800):
        return True

    for scope in iter_scopes(page):
        menu = first_existing([step4_row_locator(scope, row_index, "")])
        if not menu:
            continue

        hidden_select = first_existing([step4_row_locator(scope, row_index, "_input")])
        if hidden_select and try_select_hidden_select(hidden_select, option_texts):
            page.wait_for_timeout(700)
            if wait_for_step4_row_value(page, row_index, option_texts, timeout_ms=5000):
                return True

        trigger = first_visible(
            [
                menu.locator(".ui-selectonemenu-trigger"),
                menu.locator("label.ui-selectonemenu-label"),
                menu,
            ]
        ) or menu

        try:
            trigger.click(timeout=5000)
            page.wait_for_timeout(700)
        except Exception:
            continue

        best_item = None
        best_score = 0
        for panel_scope in iter_scopes(page):
            try:
                items = panel_scope.locator(
                    ".ui-selectonemenu-panel:visible li.ui-selectonemenu-item, li.ui-selectonemenu-item"
                )
                count = items.count()
            except Exception:
                continue

            for index in range(count):
                item = items.nth(index)
                try:
                    if not item.is_visible():
                        continue
                    item_text = item.inner_text().strip()
                except Exception:
                    continue

                if text_matches(item_text, ["Selecciona", "Select"]):
                    continue

                score = text_match_score(item_text, option_texts)
                if score > best_score:
                    best_item = item
                    best_score = score

        if not best_item:
            continue

        try:
            best_item.click(timeout=5000)
            page.wait_for_timeout(700)
        except Exception:
            continue

        if wait_for_step4_row_value(page, row_index, option_texts, timeout_ms=5000):
            return True

    if required:
        current_value = current_step4_row_text(page, row_index) or "(blank)"
        raise RuntimeError(
            f"Could not lock Step 4 row {row_index} to {option_texts[0]!r}. Current value is {current_value!r}."
        )
    return False


def select_step4_option_by_labels(page, labels, value, required=False, field_kind=None):
    row_index = wait_for_step4_row_index(page, labels, timeout_ms=2000)
    if row_index is not None:
        if select_step4_row_option(page, row_index, value, required=required, field_kind=field_kind):
            return True

    return select_option_field(page, labels, value, required=required, field_kind=field_kind)


def fill_university_diploma_specific_values_by_row(page, data, delivery_method):
    scope_value = resolve_scope(data)

    select_step4_row_option(
        page,
        2,
        data.get("training_type"),
        required=True,
        field_kind="training_type",
    )
    page.wait_for_timeout(1200)
    if scope_value:
        select_step4_row_option(
            page,
            0,
            scope_value,
            required=True,
            field_kind="scope",
        )
    if delivery_method:
        select_step4_row_option(
            page,
            1,
            delivery_method,
            required=True,
            field_kind="delivery_method",
        )

    credits_or_hours = resolve_credits_or_hours(data) or "Horas"
    selected_units = select_step4_row_option(
        page,
        3,
        credits_or_hours,
        field_kind="credits_or_hours",
    )
    if not selected_units and normalize_text(credits_or_hours) != normalize_text("Horas"):
        credits_or_hours = resolve_credits_or_hours(data) or "Horas"
        select_step4_row_option(
            page,
            3,
            credits_or_hours,
            required=True,
            field_kind="credits_or_hours",
        )
    elif not selected_units:
        select_step4_row_option(
            page,
            3,
            credits_or_hours,
            required=True,
            field_kind="credits_or_hours",
        )

    page.wait_for_timeout(1200)

    if scope_value:
        select_step4_row_option(
            page,
            0,
            scope_value,
            required=True,
            field_kind="scope",
        )
    if delivery_method:
        select_step4_row_option(
            page,
            1,
            delivery_method,
            required=True,
            field_kind="delivery_method",
        )

    credits = data.get("credits")
    hours = data.get("hours")
    hours_value = hours if hours not in (None, "", 0) else credits
    normalized_units = normalize_text(credits_or_hours)

    if normalized_units.startswith("hora") and hours_value not in (None, "", 0):
        filled_hours = fill_text_field(page, ["Horas", "Hours"], hours_value)
        if not filled_hours and credits not in (None, "", 0):
            fill_text_field(page, ["CrÃ©ditos", "Creditos", "Credits"], credits)
    elif normalized_units.startswith("credit") and credits not in (None, "", 0):
        fill_text_field(page, ["CrÃ©ditos", "Creditos", "Credits"], credits)
    else:
        if hours_value not in (None, "", 0):
            fill_text_field(page, ["Horas", "Hours"], hours_value)
        if credits not in (None, "", 0):
            fill_text_field(page, ["CrÃ©ditos", "Creditos", "Credits"], credits)

    return True


def fill_standard_specific_values_by_row(page, data):
    if not step4_row_records(page):
        return False

    training_type = resolve_formacion_continuada_training_type(data) or data.get("training_type")
    select_step4_option_by_labels(
        page,
        ["Tipo Formación", "Tipo Formacion", "Training type"],
        training_type,
        required=True,
        field_kind="training_type",
    )

    scope_value = resolve_scope(data)
    if scope_value:
        select_step4_option_by_labels(
            page,
            ["Ámbito", "Ambito", "Scope"],
            scope_value,
            field_kind="scope",
        )

    delivery_method = resolve_delivery_method(data)
    if delivery_method:
        select_step4_option_by_labels(
            page,
            [
                "Metodología de Impartición",
                "Metodologia de Imparticion",
                "Metodología",
                "Metodologia",
                "Delivery method",
            ],
            delivery_method,
            field_kind="delivery_method",
        )

    accredited_choice = resolve_accredited_choice(data)
    if accredited_choice:
        select_step4_option_by_labels(
            page,
            ["Acreditada Recibida", "Accredited received", "Acreditada"],
            accredited_choice,
            field_kind="accredited_choice",
        )
        page.wait_for_timeout(1200)

        accredited_received = resolve_accrediting_body(data)
        if accredited_received:
            select_step4_option_by_labels(
                page,
                [
                    "Órgano Acreditador Formación Recibida",
                    "Organo Acreditador Formacion Recibida",
                    "Órgano Acreditador",
                    "Organo Acreditador",
                    "Accrediting body",
                ],
                accredited_received,
                field_kind="accredited_received",
            )

    credits_or_hours = resolve_credits_or_hours(data)
    if credits_or_hours:
        select_step4_option_by_labels(
            page,
            ["Créditos/Horas", "Creditos/Horas", "Credits/Hours"],
            credits_or_hours,
            field_kind="credits_or_hours",
        )

    return True


def step4_row_input_locators(scope, row_key):
    return [
        scope.locator(f"xpath=//*[@id='form:tblFields:{row_key}:idInputText']"),
        scope.locator(f"xpath=//*[@id='form:tblFields:{row_key}:idInputTextarea']"),
    ]


def step4_row_input_exists(page, row_key):
    for scope in iter_scopes(page):
        if first_existing(step4_row_input_locators(scope, row_key)):
            return True

    return False


def current_step4_row_input_value(page, row_key):
    for scope in iter_scopes(page):
        locator = first_existing(step4_row_input_locators(scope, row_key))
        if not locator:
            continue

        try:
            return locator.input_value().strip()
        except Exception:
            pass

        try:
            return locator.inner_text().strip()
        except Exception:
            continue

    return ""


def wait_for_step4_row_input_value(page, row_key, expected_value, timeout_ms=5000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        current_value = current_step4_row_input_value(page, row_key)
        if input_value_matches(current_value, expected_value):
            return True
        if not safe_wait_for_timeout(page, 250):
            return False

    current_value = current_step4_row_input_value(page, row_key)
    return input_value_matches(current_value, expected_value)


def fill_step4_row_input(page, row_key, value, required=False):
    if value in (None, ""):
        if required:
            raise RuntimeError(f"Required Step 4 row {row_key} input is missing.")
        return False

    value_text = str(value).strip()
    if not value_text:
        if required:
            raise RuntimeError(f"Required Step 4 row {row_key} input is missing.")
        return False

    if wait_for_step4_row_input_value(page, row_key, value_text, timeout_ms=800):
        return True

    for scope in iter_scopes(page):
        locator = first_existing(step4_row_input_locators(scope, row_key))
        if not locator:
            continue

        if try_fill_locator(locator, value_text):
            try:
                locator.dispatch_event("input")
            except Exception:
                pass
            try:
                locator.dispatch_event("change")
            except Exception:
                pass
            try:
                locator.press("Tab")
            except Exception:
                pass
            page.wait_for_timeout(700)
            if wait_for_step4_row_input_value(page, row_key, value_text, timeout_ms=5000):
                return True

    if required:
        current_value = current_step4_row_input_value(page, row_key) or "(blank)"
        raise RuntimeError(
            f"Could not lock Step 4 row {row_key} to input {value_text!r}. Current value is {current_value!r}."
        )
    return False


def fill_step4_input_by_labels(page, labels, value, required=False):
    row_index = wait_for_step4_row_index(page, labels, timeout_ms=2000)
    if row_index is not None:
        if fill_step4_row_input(page, row_index, value, required=required):
            return True

    return fill_verified_text_field(page, labels, value, required=required)


def fill_step4_numeric_values(page, data, use_row_layout=False):
    credit_labels = [
        "NÂº de CrÃ©ditos",
        "NÂ° de CrÃ©ditos",
        "No de CrÃ©ditos",
        "No de Creditos",
        "CrÃ©ditos",
        "Creditos",
        "Credits",
    ]
    hour_labels = ["Horas", "Hours"]
    desired_units = normalize_text(resolve_credits_or_hours(data))
    field_values = []

    credits = data.get("credits")
    hours = data.get("hours")

    if desired_units.startswith("credit"):
        if credits not in (None, "", 0):
            field_values.append((credit_labels, credits))
        if hours not in (None, "", 0):
            field_values.append((hour_labels, hours))
    elif desired_units.startswith("hora"):
        if hours not in (None, "", 0):
            field_values.append((hour_labels, hours))
        if credits not in (None, "", 0):
            field_values.append((credit_labels, credits))
    else:
        if credits not in (None, "", 0):
            field_values.append((credit_labels, credits))
        if hours not in (None, "", 0):
            field_values.append((hour_labels, hours))

    filled_any = False
    for labels, value in field_values:
        if use_row_layout:
            filled = fill_step4_input_by_labels(page, labels, value)
        else:
            filled = fill_verified_text_field(page, labels, value)
        filled_any = bool(filled or filled_any)

    return filled_any


STEP4_SCOPE_LABELS = ["Ambito", "Scope"]
STEP4_TRAINING_TYPE_LABELS = [
    "Tipo Formacion",
    "Training type",
]
STEP4_UNIVERSITY_TRAINING_TYPE_LABELS = [
    "Tipo Titulo Propio/Diploma",
    "Tipo Titulo Propio",
]
STEP4_DELIVERY_METHOD_LABELS = [
    "Metodologia de Imparticion",
    "Metodologia",
    "Delivery method",
]
STEP4_CREDITS_OR_HOURS_LABELS = ["Creditos/Horas", "Credits/Hours"]
STEP4_COURSE_CODE_LABELS = ["Codigo Curso", "Codigo curso", "Course code"]
STEP4_CREDIT_VALUE_LABELS = [
    "N de Creditos",
    "No de Creditos",
    "Creditos",
    "Credits",
]
STEP4_HOUR_VALUE_LABELS = ["Horas", "Hours"]


def step4_option_field_exists(page, labels):
    row_index = find_step4_row_index(page, labels)
    if row_index is not None and step4_row_menu_exists(page, row_index):
        return True

    for scope in iter_scopes(page):
        if first_existing(select_option_locators(scope, labels)):
            return True

    return False


def step4_input_field_exists(page, labels):
    row_index = find_step4_row_index(page, labels)
    if row_index is not None and step4_row_input_exists(page, row_index):
        return True

    for scope in iter_scopes(page):
        if first_existing(field_locators(scope, labels)):
            return True

    return False


def finalize_step4_values(
    page,
    data,
    use_row_layout=False,
    training_type=None,
    training_type_labels=None,
    delivery_method=None,
    scope_value=None,
):
    training_type_value = training_type or resolve_formacion_continuada_training_type(data) or data.get("training_type")
    training_type_labels = training_type_labels or STEP4_TRAINING_TYPE_LABELS
    scope_choice = scope_value or resolve_scope(data)
    delivery_method_value = delivery_method or resolve_delivery_method(data) or data.get("delivery_method")
    accredited_choice = resolve_accredited_choice(data)
    accredited_received = resolve_accrediting_body(data)
    credits_or_hours = resolve_credits_or_hours(data)
    course_code = data.get("course_code")
    needs_numeric = data.get("hours") not in (None, "", 0) or data.get("credits") not in (None, "", 0)
    course_code_required = bool(course_code) and step4_input_field_exists(page, STEP4_COURSE_CODE_LABELS)

    if training_type_value and step4_option_field_exists(page, training_type_labels):
        select_step4_option_by_labels(
            page,
            training_type_labels,
            training_type_value,
            required=True,
            field_kind="training_type",
        )

    if scope_choice and step4_option_field_exists(page, STEP4_SCOPE_LABELS):
        select_step4_option_by_labels(
            page,
            STEP4_SCOPE_LABELS,
            scope_choice,
            field_kind="scope",
        )

    if delivery_method_value and step4_option_field_exists(page, STEP4_DELIVERY_METHOD_LABELS):
        select_step4_option_by_labels(
            page,
            STEP4_DELIVERY_METHOD_LABELS,
            delivery_method_value,
            field_kind="delivery_method",
        )

    if accredited_choice and step4_option_field_exists(page, ACCREDITED_CHOICE_LABELS):
        select_accredited_choice(
            page,
            accredited_choice,
            required=accredited_choice == "si",
        )
        page.wait_for_timeout(800)

    if (
        accredited_choice == "si"
        and accredited_received
        and step4_option_field_exists(page, ACCREDITING_BODY_LABELS)
    ):
        select_accrediting_body(page, accredited_received, required=True)

    if credits_or_hours and step4_option_field_exists(page, STEP4_CREDITS_OR_HOURS_LABELS):
        select_step4_option_by_labels(
            page,
            STEP4_CREDITS_OR_HOURS_LABELS,
            credits_or_hours,
            required=needs_numeric,
            field_kind="credits_or_hours",
        )
        page.wait_for_timeout(800)

    if needs_numeric:
        if not try_reveal_numeric_fields(page, data):
            save_debug(page, "step4_numeric_fields_not_visible")
            raise RuntimeError(
                "Step 4 needs credit/hour values, but the numeric inputs did not appear after setting the dropdowns."
            )
        page.wait_for_timeout(800)

    fill_step4_numeric_values(page, data, use_row_layout=use_row_layout)
    fill_step4_input_by_labels(
        page,
        STEP4_COURSE_CODE_LABELS,
        course_code,
        required=course_code_required,
    )

    if accredited_choice and step4_option_field_exists(page, ACCREDITED_CHOICE_LABELS):
        select_accredited_choice(
            page,
            accredited_choice,
            required=accredited_choice == "si",
        )
        page.wait_for_timeout(500)

    if (
        accredited_choice == "si"
        and accredited_received
        and step4_option_field_exists(page, ACCREDITING_BODY_LABELS)
    ):
        select_accrediting_body(page, accredited_received, required=True)

    if credits_or_hours and step4_option_field_exists(page, STEP4_CREDITS_OR_HOURS_LABELS):
        select_step4_option_by_labels(
            page,
            STEP4_CREDITS_OR_HOURS_LABELS,
            credits_or_hours,
            required=needs_numeric,
            field_kind="credits_or_hours",
        )
        page.wait_for_timeout(500)

    if needs_numeric:
        fill_step4_numeric_values(page, data, use_row_layout=use_row_layout)

    if data.get("credits") not in (None, "", 0) and step4_input_field_exists(page, STEP4_CREDIT_VALUE_LABELS):
        fill_step4_input_by_labels(
            page,
            STEP4_CREDIT_VALUE_LABELS,
            data.get("credits"),
            required=True,
        )

    if data.get("hours") not in (None, "", 0) and step4_input_field_exists(page, STEP4_HOUR_VALUE_LABELS):
        fill_step4_input_by_labels(
            page,
            STEP4_HOUR_VALUE_LABELS,
            data.get("hours"),
            required=True,
        )

    fill_step4_input_by_labels(
        page,
        STEP4_COURSE_CODE_LABELS,
        course_code,
        required=course_code_required,
    )

    return True


def select_option_candidates(value, field_kind=None):
    candidates = []
    if value not in (None, ""):
        raw_value = str(value)
        candidates.append(raw_value)
        repaired_value = repair_mojibake_text(raw_value)
        if normalize_text(repaired_value) != normalize_text(raw_value):
            candidates.append(repaired_value)

    normalized = normalize_text(value)

    training_mapping = {
        "congreso": ["Congreso", "Congress", "Conference"],
        "congress": ["Congreso", "Congress", "Conference"],
        "conference": ["Congreso", "Congress", "Conference"],
        "curso": ["Curso", "Course", "Training"],
        "course": ["Curso", "Course", "Training"],
        "training": ["Curso", "Course", "Training"],
        "certificate": ["Curso", "Course", "Training"],
        "certificado": ["Curso", "Course", "Training"],
        "diploma de especializacion": ["Diploma de Especializacion", "Diploma de especializacion", "Specialization Diploma"],
        "specialization diploma": ["Diploma de Especializacion", "Diploma de especializacion", "Specialization Diploma"],
        "specialisation diploma": ["Diploma de Especializacion", "Diploma de especializacion", "Specialization Diploma"],
        "jornada": ["Jornada", "Conference Day"],
        "master": ["Master Titulo Propio", "Master Universitario No Grado Academico", "Professional Master"],
        "master profesional": ["Master profesional", "Professional Master"],
        "professional master": ["Master profesional", "Professional Master"],
        "master titulo propio": ["Master Titulo Propio", "Master Own Degree"],
        "master universitario no grado academico": ["Master Universitario No Grado Academico", "University Master Without Academic Degree"],
        "master de formacion permanente": ["Master Universitario No Grado Academico", "Master Titulo Propio"],
        "curso universitario": ["Curso Universitario", "Curso", "Cursos, Diplomas o Certificaciones de Extension Universitaria", "University Course"],
        "certificado universitario": ["Cursos, Diplomas o Certificaciones de Extension Universitaria", "University Certificate"],
        "diploma universitario": ["Cursos, Diplomas o Certificaciones de Extension Universitaria", "University Diploma"],
        "cursos, diplomas o certificaciones de extension universitaria": ["Cursos, Diplomas o Certificaciones de Extension Universitaria", "University Extension Courses, Diplomas or Certifications"],
        "titulo propio diploma de especializacion": ["Titulo Propio Diploma de Especializacion", "Own Degree Specialization Diploma"],
        "titulo propio experto universitario": ["Titulo Propio Experto Universitario", "Own Degree University Expert"],
        "titulo propio especialista universitario": ["Titulo Propio Especialista Universitario", "Own Degree University Specialist"],
        "titulo propio universitario distinto a los anteriores": ["Titulo Propio Universitario distinto a los Anteriores", "Other University Own Degree"],
        "experto universitario": ["Experto Universitario", "Titulo Propio Experto Universitario", "University Expert"],
        "seminario": ["Seminario", "Seminar", "Webinar"],
        "seminar": ["Seminario", "Seminar", "Webinar"],
        "webinar": ["Seminario", "Seminar", "Webinar"],
        "sesion clinica": ["Sesion Clinica", "Sesión Clínica", "Clinical Session"],
        "clinical session": ["Sesion Clinica", "Sesión Clínica", "Clinical Session"],
        "taller": ["Taller", "Workshop"],
        "workshop": ["Taller", "Workshop"],
        "otros": ["Otros", "Other"],
        "other": ["Otros", "Other"],
        "titulo propio": ["TÃ­tulo Propio", "Titulo Propio", "Own Degree"],
        "estancia formativa": ["Estancia Formativa", "Training Placement"],
    }
    if field_kind == "training_type" and normalized in training_mapping:
        candidates.extend(training_mapping[normalized])
    if field_kind == "training_type" and normalized == "titulo propio":
        candidates.append("Titulo Propio Universitario distinto a los Anteriores")

    if field_kind == "delivery_method":
        if normalized in ("online", "on line", "distancia", "a distancia", "teleformacion"):
            candidates.extend(
                [
                    "Online",
                    "On line",
                    "No Presencial",
                    "A distancia",
                    "Teleformación",
                    "Teleformacion",
                    "Distance learning",
                    "Remote learning",
                ]
            )
        elif normalized in ("presencial",):
            candidates.extend(["Presencial", "Face-to-face", "On-site"])
        elif normalized in ("mixta", "semipresencial"):
            candidates.extend(["Mixta", "Semipresencial", "Blended", "Hybrid"])

    if field_kind == "delivery_method":
        if any(token in normalized for token in ("om)", "mooc", "online masivos", "abiertos")):
            candidates.extend(["Cursos online masivos y abiertos (OM)", "Massive open online courses (OM)"])
        if any(
            token in normalized
            for token in ("e-learning", "elearning", "online", "on line", "virtual", "modalidad e-learning")
        ):
            candidates.extend(["A distancia modalidad e-learning (V)", "Distance e-learning mode (V)"])
        if normalized in ("presencial", "formacion presencial", "formacion presencial (p)"):
            candidates.extend(["Formacion Presencial (P)", "Face-to-face Training (P)", "On-site Training (P)"])
        if normalized in ("mixta", "semipresencial", "semipresencial (s)"):
            candidates.extend(["Semipresencial (S)", "Blended learning (S)", "Hybrid learning (S)"])
        if normalized in ("distancia", "a distancia", "a distancia (d)"):
            candidates.extend(["A distancia (D)", "Distance learning (D)"])

    if field_kind == "scope":
        if normalized in ("autonomico", "regional", "autonomic"):
            candidates.extend(["Autonomico", "Autonómico", "Regional", "Autonomic"])
        elif normalized in ("nacional", "espana", "spain", "national"):
            candidates.extend(["Nacional", "National"])
        elif normalized in ("comunitario", "union europea", "ue", "eu", "community", "european union"):
            candidates.extend(["Comunitario", "Community", "European Union", "EU"])
        elif normalized in ("extracomunitario", "internacional", "extra ue", "extra-eu", "international", "outside eu", "non-eu"):
            candidates.extend(["Extracomunitario", "International", "Outside EU", "Non-EU"])

    if field_kind == "credits_or_hours":
        if normalized.startswith("credit") or normalized == "ects":
            candidates.extend(["Creditos", "Créditos", "Credits"])
        elif normalized.startswith("hora") or normalized == "hours":
            candidates.extend(["Horas", "Hours"])

    if field_kind in ("yes_no", "accredited_choice"):
        if normalized in ("si", "sÃ­", "sí", "yes", "true", "1"):
            candidates.extend(["Si", "SÃ­", "Sí", "Yes"])
        elif normalized in ("no", "false", "0"):
            candidates.extend(["No"])

    if field_kind == "accredited_received":
        if any(token in normalized for token in ("secretaria general de salud publica", "autonomica", "junta de andalucia", "i+d+i")):
            candidates.extend(
                [
                    "Acreditado por la Comision de Formacion Continuada Autonomica",
                    "Accredited by the Regional Continuing Education Commission",
                ]
            )
        if any(token in normalized for token in ("sistema nacional de salud", "del sns", "sns")):
            candidates.extend(
                [
                    "Acreditado por la Comision de Formacion Continuada del SNS",
                    "Accredited by the National Health System Continuing Education Commission",
                ]
            )
        if any(token in normalized for token in ("consejo internacional de enfermeria", "international council of nurses")):
            candidates.extend(
                [
                    "Acreditado por el Consejo Internacional de Enfermeria",
                    "Accredited by the International Council of Nurses",
                ]
            )
        if "eaccme" in normalized or "european accreditation council" in normalized:
            candidates.extend(
                [
                    "Acreditado por European Accreditation Council For CME (EACCME)",
                    "Accredited by European Accreditation Council For CME (EACCME)",
                ]
            )
        if "american medical association" in normalized or re.search(r"\bama\b", normalized):
            candidates.extend(
                [
                    "Acreditado por American Medical Association (AMA)",
                    "Acreditado por  American Medical Association (AMA)",
                    "Accredited by American Medical Association (AMA)",
                ]
            )
        if "royal college of physicians and surgeons of canada" in normalized:
            candidates.extend(
                [
                    "Acreditado por Royal College of Physicians and Surgeons of Canada",
                    "Accredited by Royal College of Physicians and Surgeons of Canada",
                ]
            )

    return dedupe_preserve_order(candidates)


def field_is_visible(page, labels):
    for scope in iter_scopes(page):
        if first_visible(field_locators(scope, labels)):
            return True
    return False


def set_checkbox_state(page, labels, checked):
    target_state = bool(checked)
    for label in expand_ui_text_variants(labels):
        for scope in iter_scopes(page):
            locator = first_visible(
                [
                    scope.get_by_label(label, exact=False),
                    scope.locator(
                        f"xpath=(//label[contains(normalize-space(.), \"{label}\")]/preceding::*[self::input[@type='checkbox']][1])[1]"
                    ),
                    scope.locator(
                        f"xpath=(//*[contains(normalize-space(.), \"{label}\")]/ancestor::*[self::div or self::td or self::tr][1]//input[@type='checkbox'][1])[1]"
                    ),
                ]
            )
            if not locator:
                continue

            try:
                current = locator.is_checked()
            except Exception:
                current = None

            if current == target_state:
                return True

            try:
                locator.set_checked(target_state)
                page.wait_for_timeout(300)
                return True
            except Exception:
                pass

            try:
                locator.click(timeout=3000)
                page.wait_for_timeout(300)
                return True
            except Exception:
                continue

    return not checked


def resolve_accredited_choice(data):
    explicit = normalize_text(data.get("accredited_choice"))
    if explicit in ("si", "no"):
        return explicit

    accredited_body = normalize_text(data.get("accredited_received"))
    if accredited_body or data.get("credits") not in (None, "", 0):
        return "si"

    return ""


def evidence_text(data, *keys):
    texts = []
    field_evidence = data.get("field_evidence")
    if not isinstance(field_evidence, dict):
        field_evidence = {}

    for key in keys:
        value = data.get(key)
        if value not in (None, ""):
            texts.append(str(value))

        evidence_value = field_evidence.get(key)
        if evidence_value not in (None, ""):
            texts.append(str(evidence_value))

    return normalize_text(" ".join(texts))


def resolve_formacion_continuada_training_type(data):
    explicit = normalize_text(data.get("training_type"))
    if explicit:
        mapping = {
            "congreso": "Congreso",
            "congress": "Congreso",
            "conference": "Congreso",
            "curso": "Curso",
            "course": "Curso",
            "training": "Curso",
            "certificate": "Curso",
            "certificado": "Curso",
            "diploma de especializacion": "Diploma de especialización",
            "specialization diploma": "Diploma de especialización",
            "specialisation diploma": "Diploma de especialización",
            "jornada": "Jornada",
            "master profesional": "Master profesional",
            "professional master": "Master profesional",
            "master": "Master profesional",
            "seminario": "Seminario",
            "seminar": "Seminario",
            "webinar": "Seminario",
            "sesion clinica": "Sesión Clínica",
            "clinical session": "Sesión Clínica",
            "taller": "Taller",
            "workshop": "Taller",
            "otros": "Otros",
            "other": "Otros",
        }
        if explicit in mapping:
            return mapping[explicit]
        return data.get("training_type")

    combined = evidence_text(
        data,
        "training_type",
        "merit_name",
        "document_description",
        "reasoning_summary",
    )
    for token, option in (
        ("sesion clinica", "Sesión Clínica"),
        ("clinical session", "Sesión Clínica"),
        ("congreso", "Congreso"),
        ("congress", "Congreso"),
        ("conference", "Congreso"),
        ("simposio", "Congreso"),
        ("symposium", "Congreso"),
        ("seminario", "Seminario"),
        ("seminar", "Seminario"),
        ("webinar", "Seminario"),
        ("taller", "Taller"),
        ("workshop", "Taller"),
        ("jornada", "Jornada"),
        ("diploma de especializacion", "Diploma de especialización"),
        ("specialization diploma", "Diploma de especialización"),
        ("specialisation diploma", "Diploma de especialización"),
        ("master", "Master profesional"),
        ("curso", "Curso"),
        ("course", "Curso"),
        ("training", "Curso"),
        ("certificate", "Curso"),
        ("certificado", "Curso"),
    ):
        if token in combined:
            return option

    return ""


def resolve_accrediting_body(data):
    explicit = normalize_text(data.get("accredited_received"))
    if explicit:
        if any(token in explicit for token in ("secretaria general de salud publica", "autonomica", "junta de andalucia", "i+d+i")):
            return "Acreditado por la Comisión de Formación Continuada Autonómica"
        if any(token in explicit for token in ("sistema nacional de salud", "del sns", "sns")):
            return "Acreditado por la Comisión de Formación Continuada del SNS"
        if any(token in explicit for token in ("consejo internacional de enfermeria", "international council of nurses")):
            return "Acreditado por el Consejo Internacional de Enfermería"
        if "eaccme" in explicit or "european accreditation council" in explicit:
            return "Acreditado por European Accreditation Council For CME (EACCME)"
        if "american medical association" in explicit or re.search(r"\bama\b", explicit):
            return "Acreditado por  American Medical Association (AMA)"
        if "royal college of physicians and surgeons of canada" in explicit:
            return "Acreditado por Royal College of Physicians and Surgeons of Canada"
        return data.get("accredited_received")

    combined = evidence_text(
        data,
        "accredited_received",
        "accredited_choice",
        "reasoning_summary",
        "document_description",
    )
    if any(token in combined for token in ("secretaria general de salud publica", "autonomica", "junta de andalucia", "i+d+i")):
        return "Acreditado por la Comisión de Formación Continuada Autonómica"
    if any(token in combined for token in ("sistema nacional de salud", "del sns", "sns")):
        return "Acreditado por la Comisión de Formación Continuada del SNS"
    if any(token in combined for token in ("consejo internacional de enfermeria", "international council of nurses")):
        return "Acreditado por el Consejo Internacional de Enfermería"
    if "eaccme" in combined or "european accreditation council" in combined:
        return "Acreditado por European Accreditation Council For CME (EACCME)"
    if "american medical association" in combined or re.search(r"\bama\b", combined):
        return "Acreditado por  American Medical Association (AMA)"
    if "royal college of physicians and surgeons of canada" in combined:
        return "Acreditado por Royal College of Physicians and Surgeons of Canada"

    return ""


ACCREDITED_CHOICE_LABELS = ["Acreditada Recibida", "Accredited received", "Acreditada"]
ACCREDITING_BODY_LABELS = [
    "Ã“rgano Acreditador FormaciÃ³n Recibida",
    "Organo Acreditador Formacion Recibida",
    "Ã“rgano Acreditador",
    "Organo Acreditador",
    "Accrediting body",
]


ACCREDITED_CHOICE_LABELS = ["Acreditada Recibida", "Accredited received", "Acreditada"]
ACCREDITING_BODY_LABELS = [
    "Organo Acreditador Formacion Recibida",
    "Organo Acreditador",
    "Accrediting body",
]


def select_accredited_choice(page, choice, required=False):
    return select_step4_option_by_labels(
        page,
        ACCREDITED_CHOICE_LABELS,
        choice,
        required=required,
        field_kind="accredited_choice",
    )


def select_accrediting_body(page, accrediting_body, required=False):
    return select_step4_option_by_labels(
        page,
        ACCREDITING_BODY_LABELS,
        accrediting_body,
        required=required,
        field_kind="accredited_received",
    )


def apply_accreditation_fields(page, data, required=False):
    accredited_choice = resolve_accredited_choice(data)
    if not accredited_choice:
        return False

    choice_selected = select_accredited_choice(
        page,
        accredited_choice,
        required=required or accredited_choice == "si",
    )
    if not choice_selected:
        return False

    page.wait_for_timeout(1200)

    if accredited_choice != "si":
        return True

    accredited_received = resolve_accrediting_body(data)
    if accredited_received:
        select_accrediting_body(page, accredited_received, required=required)

    return True


def resolve_credits_or_hours(data):
    explicit = normalize_text(data.get("credits_or_hours"))
    if explicit.startswith("credit") or explicit == "ects":
        return "Créditos"
    if explicit.startswith("hora") or explicit == "hours":
        return "Horas"

    if data.get("credits") not in (None, "", 0):
        return "Créditos"
    if data.get("hours") not in (None, "", 0):
        return "Horas"

    return ""


def resolve_delivery_method(data):
    explicit = normalize_text(data.get("delivery_method"))
    if explicit:
        if any(token in explicit for token in ("om)", "mooc", "online masivos", "abiertos")):
            return "Cursos online masivos y abiertos (OM)"
        if any(
            token in explicit
            for token in ("e-learning", "elearning", "online", "on line", "virtual", "modalidad e-learning")
        ):
            return "A distancia modalidad e-learning (V)"
        if explicit in ("presencial", "formacion presencial", "formacion presencial (p)"):
            return "FormaciÃ³n Presencial (P)"
        if explicit in ("mixta", "semipresencial", "semipresencial (s)"):
            return "Semipresencial (S)"
        if explicit in ("distancia", "a distancia", "a distancia (d)"):
            return "A distancia (D)"
        return data.get("delivery_method")

    return ""


def resolve_scope(data):
    explicit = normalize_text(data.get("scope"))
    if explicit in ("autonomico", "nacional", "comunitario", "extracomunitario"):
        return {
            "autonomico": "Autonómico",
            "nacional": "Nacional",
            "comunitario": "Comunitario",
            "extracomunitario": "Extracomunitario",
        }[explicit]

    return ""


def try_reveal_numeric_fields(page, data):
    needs_numeric = data.get("hours") not in (None, "", 0) or data.get("credits") not in (None, "", 0)
    if not needs_numeric:
        return True

    if field_is_visible(page, ["Horas", "Hours"]) or field_is_visible(page, ["CrÃ©ditos", "Creditos", "Credits"]):
        return True

    credits_or_hours = resolve_credits_or_hours(data)
    if credits_or_hours:
        if select_step4_option_by_labels(
            page,
            ["CrÃ©ditos/Horas", "Creditos/Horas", "Credits/Hours"],
            credits_or_hours,
            field_kind="credits_or_hours",
        ) or select_option_field(
            page,
            ["CrÃ©ditos/Horas", "Creditos/Horas", "Credits/Hours"],
            credits_or_hours,
            field_kind="credits_or_hours",
        ):
            page.wait_for_timeout(1200)
            if field_is_visible(page, ["Horas", "Hours"]) or field_is_visible(
                page,
                ["CrÃƒÂ©ditos", "Creditos", "Credits"],
            ):
                return True

    desired_choice = resolve_accredited_choice(data)
    if desired_choice in ("si", "no"):
        choices = [desired_choice]
    else:
        choices = ["si", "no"]
    for choice in choices:
        if not choice:
            continue

        row_selection_applied = select_step4_option_by_labels(
            page,
            ["Acreditada Recibida", "Accredited received", "Acreditada"],
            choice,
            field_kind="accredited_choice",
        )

        if row_selection_applied or select_option_field(
            page,
            ["Acreditada Recibida", "Accredited received", "Acreditada"],
            choice,
            field_kind="accredited_choice",
        ):
            page.wait_for_timeout(1200)
            if field_is_visible(page, ["Horas", "Hours"]) or field_is_visible(page, ["CrÃ©ditos", "Creditos", "Credits"]):
                return True

    return field_is_visible(page, ["Horas", "Hours"]) or field_is_visible(page, ["CrÃ©ditos", "Creditos", "Credits"])


def _unused_try_reveal_numeric_fields_override(page, data):
    needs_numeric = data.get("hours") not in (None, "", 0) or data.get("credits") not in (None, "", 0)
    if not needs_numeric:
        return True

    if field_is_visible(page, ["Horas", "Hours"]) or field_is_visible(page, ["CrÃƒÂ©ditos", "Creditos", "Credits"]):
        return True

    desired_choice = resolve_accredited_choice(data)
    if desired_choice == "si":
        choices = ["si"]
    else:
        # Numeric fields are usually revealed behind an affirmative accreditation choice.
        # Prefer "si" unless the data explicitly says otherwise.
        choices = dedupe_preserve_order([desired_choice, "si", "no"])

    for choice in choices:
        if not choice:
            continue

        if select_accredited_choice(
            page,
            choice,
            required=choice == "si" and desired_choice == "si",
        ):
            page.wait_for_timeout(1200)
            if field_is_visible(page, ["Horas", "Hours"]) or field_is_visible(
                page,
                ["CrÃƒÂ©ditos", "Creditos", "Credits"],
            ):
                return True

    return field_is_visible(page, ["Horas", "Hours"]) or field_is_visible(
        page,
        ["CrÃƒÂ©ditos", "Creditos", "Credits"],
    )


def step_specific_values_university_diploma(page, data):
    delivery_method = resolve_delivery_method(data)
    if step4_row_records(page):
        if fill_university_diploma_specific_values_by_row(page, data, delivery_method):
            return True

    select_option_field(
        page,
        ["Ámbito", "Ambito", "Scope"],
        resolve_scope(data),
        field_kind="scope",
    )
    select_option_field(
        page,
        [
            "Metodología de Impartición",
            "Metodologia de Imparticion",
            "Metodología",
            "Metodologia",
            "Delivery method",
        ],
        delivery_method,
        required=True,
        field_kind="delivery_method",
    )
    select_option_field(
        page,
        [
            "Tipo Título Propio/Diploma",
            "Tipo Titulo Propio/Diploma",
            "Tipo Título Propio",
            "Tipo Titulo Propio",
        ],
        data.get("training_type"),
        required=True,
        field_kind="training_type",
    )
    select_option_field(
        page,
        ["Ãmbito", "Ambito", "Scope"],
        resolve_scope(data),
        required=True,
        field_kind="scope",
    )
    select_option_field(
        page,
        [
            "MetodologÃ­a de ImparticiÃ³n",
            "Metodologia de Imparticion",
            "MetodologÃ­a",
            "Metodologia",
            "Delivery method",
        ],
        delivery_method,
        required=True,
        field_kind="delivery_method",
    )

    # Prefer the classified unit from the JSON/text extraction and only fall
    # back to "Horas" if that exact option is not available on the page.
    credits_or_hours = resolve_credits_or_hours(data) or "Horas"
    selected_units = select_option_field(
        page,
        ["Créditos/Horas", "Creditos/Horas", "Credits/Hours"],
        credits_or_hours,
        field_kind="credits_or_hours",
    )
    if not selected_units and normalize_text(credits_or_hours) != normalize_text("Horas"):
        fallback_units = "Horas"
        select_option_field(
            page,
            ["Créditos/Horas", "Creditos/Horas", "Credits/Hours"],
            fallback_units,
            required=True,
            field_kind="credits_or_hours",
        )
        credits_or_hours = fallback_units
    elif not selected_units:
        select_option_field(
            page,
            ["CrÃ©ditos/Horas", "Creditos/Horas", "Credits/Hours"],
            credits_or_hours,
            required=True,
            field_kind="credits_or_hours",
        )

    page.wait_for_timeout(1200)

    credits = data.get("credits")
    hours = data.get("hours")
    hours_value = hours if hours not in (None, "", 0) else credits
    normalized_units = normalize_text(credits_or_hours)

    if normalized_units.startswith("hora") and hours_value not in (None, "", 0):
        filled_hours = fill_text_field(page, ["Horas", "Hours"], hours_value)
        if not filled_hours and credits not in (None, "", 0):
            fill_text_field(page, ["Créditos", "Creditos", "Credits"], credits)
    elif normalized_units.startswith("credit") and credits not in (None, "", 0):
        fill_text_field(page, ["Créditos", "Creditos", "Credits"], credits)
    else:
        if hours_value not in (None, "", 0):
            fill_text_field(page, ["Horas", "Hours"], hours_value)
        if credits not in (None, "", 0):
            fill_text_field(page, ["Créditos", "Creditos", "Credits"], credits)


    return True


def document_table_has_rows(page):
    for scope in iter_scopes(page):
        try:
            if scope.get_by_text("No hay registros", exact=False).count() > 0:
                continue
        except Exception:
            pass

        locator = first_visible(
            [
                scope.locator("tbody[id$='data'] tr"),
                scope.locator("table tbody tr"),
            ]
        )
        if not locator:
            continue

        try:
            text = locator.inner_text().strip()
        except Exception:
            text = ""

        if text and not text_matches(text, ["No hay registros", "No records"]):
            return True

    return False


def document_upload_error_message(page):
    messages = visible_error_messages(page)
    if messages:
        return "; ".join(messages)

    for scope in iter_scopes(page):
        locator = first_visible(
            [
                scope.locator("#form\\:errorPanel .h2"),
                scope.locator("[id$='errorPanel'] .h2"),
                scope.locator(".ui-fileupload-content .ui-messages-error-summary"),
                scope.locator(".ui-fileupload-content .ui-messages-error-detail"),
            ]
        )
        if not locator:
            continue

        try:
            text = locator.inner_text().strip()
        except Exception:
            text = ""

        if text:
            return text

    return ""


def wait_for_document_row(page, timeout_ms=30000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        if document_table_has_rows(page):
            return True
        if document_upload_error_message(page):
            return False
        page.wait_for_timeout(400)

    return document_table_has_rows(page)


def click_document_continue(page):
    direct = []
    for scope in iter_scopes(page):
        direct.extend(
            [
                scope.locator("#form\\:btnFirmarDocumentos"),
                scope.locator("button[id$='btnFirmarDocumentos']"),
            ]
        )

    locator = first_visible(direct)
    if locator:
        try:
            locator.click(timeout=5000)
            page.wait_for_timeout(1200)
            return True
        except Exception:
            pass

    return click_if_found(
        page,
        ["Firmar Documentos", "Sign Documents", "Continuar", "Guardar y Siguiente", "Siguiente", "Next"],
        timeout_ms=7000,
    )


def click_step6_sign_and_finish(page):
    direct = []
    for scope in iter_scopes(page):
        direct.extend(
            [
                scope.locator("#formFirma\\:firmarButton"),
                scope.locator("button[id$='firmarButton']"),
            ]
        )

    locator = first_visible(direct)
    if locator:
        try:
            locator.click(timeout=5000)
            page.wait_for_timeout(1200)
            return True
        except Exception:
            pass

    return click_if_found_with_retry(
        page,
        ["Firmar y Terminar", "Firmar", "Sign and Finish", "Sign"],
        timeout_ms=8000,
    )


def click_step6_server_sign(page):
    direct = []
    for scope in iter_scopes(page):
        direct.extend(
            [
                scope.locator("#firmarservidorButton"),
                scope.locator("button[id$='firmarservidorButton']"),
            ]
        )

    locator = first_visible(direct)
    if locator:
        try:
            locator.click(timeout=5000)
            page.wait_for_timeout(1500)
            return True
        except Exception:
            pass

    return click_if_found_with_retry(
        page,
        ["Firmar con Servidor", "Firma con Servidor", "Sign with Server"],
        timeout_ms=10000,
    )


def step6_has_unsigned_documents(page):
    for scope in iter_scopes(page):
        try:
            if scope.get_by_text("No Firmado", exact=False).count() > 0:
                return True
        except Exception:
            continue
    return False


def wait_for_step6_completion(page, timeout_ms=30000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        accept_modal_if_present(page)

        if not wait_for_step(page, STEP6_SIGN_MARKERS, active_titles=STEP6_ACTIVE_TITLES, timeout_ms=800):
            return True

        if not step6_has_unsigned_documents(page):
            return True

        if not safe_wait_for_timeout(page, 500):
            return False

    return not step6_has_unsigned_documents(page)


def classification_warnings(data):
    required_fields = {
        "vec_route": data.get("vec_route"),
        "merit_name": data.get("merit_name"),
        "institution_name": data.get("institution_name"),
        "center_search_text": data.get("center_search_text") or data.get("institution_name"),
        "start_date": data.get("start_date"),
        "end_date": data.get("end_date") or data.get("year"),
    }
    recommended_fields = {
        "scope": data.get("scope"),
        "training_type": data.get("training_type"),
        "delivery_method": data.get("delivery_method"),
        "accredited_choice": data.get("accredited_choice"),
        "accredited_received": data.get("accredited_received"),
        "center_result_text": data.get("center_result_text"),
        "course_code": data.get("course_code"),
        "credits_or_hours": data.get("credits_or_hours"),
        "document_description": data.get("document_description"),
        "hours": data.get("hours"),
        "credits": data.get("credits"),
    }

    missing_required = [key for key, value in required_fields.items() if value in (None, "", 0)]
    missing_recommended = [key for key, value in recommended_fields.items() if value in (None, "", 0)]

    date_pattern = re.compile(r"^\d{2}/\d{2}/\d{4}$")
    bad_dates = []
    for key in ("start_date", "end_date"):
        value = data.get(key)
        normalized_value = normalize_form_date(value, end_of_year=(key == "end_date"))
        if value and not date_pattern.match(str(normalized_value)):
            bad_dates.append((key, value))

    return missing_required, missing_recommended, bad_dates


def classification_blockers(data, missing_required, missing_recommended, bad_dates):
    blockers = []
    route = str(data.get("vec_route") or "").strip()
    blocking_recommended = []

    if route in TRAINING_ROUTE_VALUES:
        for field in (
            "training_type",
            "delivery_method",
            "accredited_choice",
            "course_code",
            "credits_or_hours",
            "document_description",
        ):
            if field in missing_recommended:
                blocking_recommended.append(field)

    if route == "manual_review":
        blockers.append("vec_route=manual_review")
    if missing_required:
        blockers.append("missing required fields: " + ", ".join(missing_required))
    if blocking_recommended:
        blockers.append("missing critical training fields: " + ", ".join(blocking_recommended))
    if bad_dates:
        blockers.append(
            "bad date formats: " + ", ".join(f"{key}={value}" for key, value in bad_dates)
        )

    alerts = [str(value or "") for value in data.get("alerts") or []]
    normalized_alerts = [normalize_text(value) for value in alerts]
    blocking_alerts = [
        alert
        for alert, normalized in zip(alerts, normalized_alerts)
        if any(marker in normalized for marker in BLOCKING_CLASSIFICATION_ALERT_MARKERS)
    ]
    if blocking_alerts:
        blockers.append("blocking alerts: " + " | ".join(blocking_alerts))

    return blockers


def print_classification_summary(data, missing_required, missing_recommended, bad_dates, blockers):
    print("Classification loaded:")
    print(json.dumps(data, indent=2, ensure_ascii=False))
    print(f"Resolved PDF: {data['resolved_pdf_path']}")
    if missing_required:
        print(f"Missing required data: {', '.join(missing_required)}")
    else:
        print("Required classification fields look complete.")
    if missing_recommended:
        print(f"Recommended but blank fields: {', '.join(missing_recommended)}")
    if bad_dates:
        print(
            "Unexpected date formats: "
            + ", ".join(f"{key}={value}" for key, value in bad_dates)
        )
    if blockers:
        print("Blocking classification issues:")
        for blocker in blockers:
            print(f"- {blocker}")


def classify_next_ready_document(classification_backend, max_attempts=MAX_CLASSIFICATION_ATTEMPTS):
    skipped_blocked_reasons = []

    for attempt in range(1, max_attempts + 1):
        print()
        print(f"Classification pass {attempt}/{max_attempts}: selecting the next PDF.")

        try:
            classification = classifier_engine.classify_next_document(
                explicit_pdf_path=None,
                pdf_directory=PDF_DIRECTORY,
                output=CLASSIFICATION_FILE,
                text_output=CLASSIFICATION_TEXT_FILE,
                backend=classification_backend,
            )
        except Exception as exc:
            if skipped_blocked_reasons:
                raise RuntimeError(
                    "No remaining complete PDFs were available after skipping blocked classifications. "
                    f"Last blocked reasons: {' | '.join(skipped_blocked_reasons[-5:])}. "
                    f"Final classifier error: {exc}"
                ) from exc
            raise

        data = dict(classification["result"])
        data["resolved_pdf_path"] = resolve_pdf_path(data)

        missing_required, missing_recommended, bad_dates = classification_warnings(data)
        blockers = classification_blockers(data, missing_required, missing_recommended, bad_dates)
        print_classification_summary(data, missing_required, missing_recommended, bad_dates, blockers)

        if not blockers:
            return data, classification, missing_required, missing_recommended, bad_dates

        print()
        print("Blocked classification detected. Running a detailed OpenAI re-review with the official reference PDFs before any skip...")
        try:
            reviewed_data = classifier_engine.review_blocked_classification_with_openai(
                document_text=classification.get("document_text") or "",
                pdf_path=classification["pdf_path"],
                current_result=data,
                blockers=blockers,
            )
        except Exception as exc:
            raise RuntimeError(
                "A blocked document could not be re-reviewed with OpenAI before skipping. "
                f"Last error: {exc}"
            ) from exc

        reviewed_missing_required, reviewed_missing_recommended, reviewed_bad_dates = classification_warnings(reviewed_data)
        reviewed_blockers = classification_blockers(
            reviewed_data,
            reviewed_missing_required,
            reviewed_missing_recommended,
            reviewed_bad_dates,
        )
        print_classification_summary(
            reviewed_data,
            reviewed_missing_required,
            reviewed_missing_recommended,
            reviewed_bad_dates,
            reviewed_blockers,
        )

        if not reviewed_blockers:
            print("Accepted after the detailed OpenAI re-review.")
            return (
                reviewed_data,
                classification,
                reviewed_missing_required,
                reviewed_missing_recommended,
                reviewed_bad_dates,
            )

        data = reviewed_data
        missing_required = reviewed_missing_required
        missing_recommended = reviewed_missing_recommended
        bad_dates = reviewed_bad_dates
        blockers = reviewed_blockers

        reason = "Incomplete/manual-review classification blocked automated upload: " + " | ".join(blockers)
        skipped_blocked_reasons.append(reason)
        print(f"Skipping classified PDF: {reason}")
        archived_pdf_path = classifier_engine.skip_pdf_candidate(
            classification["pdf_path"],
            reason,
            status="skipped_incomplete_classification",
        )
        try:
            snapshot_path = save_skipped_classification_snapshot(
                classification,
                data,
                blockers,
                reason,
                archived_pdf_path=archived_pdf_path,
            )
            print(f"Saved skipped classification snapshot: {snapshot_path}")
        except Exception as exc:
            print(f"Warning: Could not save skipped classification snapshot: {exc}")

    raise RuntimeError(
        f"Reached the maximum classification attempts ({max_attempts}) without finding a complete PDF."
    )


def normalize_form_date(value, end_of_year=False):
    if value in (None, ""):
        return ""

    if isinstance(value, dt.datetime):
        return value.strftime("%d/%m/%Y")

    if isinstance(value, dt.date):
        return value.strftime("%d/%m/%Y")

    text = str(value).strip()
    if re.fullmatch(r"\d{4}", text):
        return f"{'31/12' if end_of_year else '01/01'}/{text}"

    trimmed_text = text.split("T", 1)[0].strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            parsed = dt.datetime.strptime(trimmed_text, fmt)
            return parsed.strftime("%d/%m/%Y")
        except ValueError:
            continue

    match = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", trimmed_text)
    if match:
        day, month, year = match.groups()
        return f"{int(day):02d}/{int(month):02d}/{year}"

    return text


def parse_form_date_parts(value):
    text = normalize_form_date(value)
    match = re.fullmatch(r"(\d{2})/(\d{2})/(\d{4})", text)
    if not match:
        return None

    day, month, year = match.groups()
    return int(year), int(month), int(day)


def resolve_step1_dates(data):
    start_date = normalize_form_date(data.get("start_date") or "", end_of_year=False)
    end_date = normalize_form_date(
        data.get("end_date") or data.get("year") or "",
        end_of_year=True,
    )

    route = str(data.get("vec_route") or "").strip()
    year_text = str(data.get("year") or "").strip()
    start_parts = parse_form_date_parts(start_date)
    end_parts = parse_form_date_parts(end_date)

    if route == "diplomas_titulos_propios_universitarios" and re.fullmatch(r"\d{4}", year_text):
        merit_year = int(year_text)
        has_inconsistent_dates = (
            not start_date
            or not end_date
            or (start_parts and start_parts[0] != merit_year)
            or (end_parts and end_parts[0] != merit_year)
            or (start_parts and end_parts and start_parts > end_parts)
        )
        if has_inconsistent_dates:
            start_date = f"01/01/{merit_year}"
            end_date = f"31/12/{merit_year}"
            print(
                "Step 1 date repair: using university diploma year window "
                f"{start_date!r} -> {end_date!r}"
            )

    return start_date, end_date


def step_dates(page, data):
    print("Step 1: Dates")

    start_date, end_date = resolve_step1_dates(data)

    print(f"Step 1 values: start={start_date!r}, end={end_date!r}")

    fill_text_field(page, ["Fecha de Inicio", "Fecha Inicio", "Start date"], start_date, required=True)
    fill_text_field(
        page,
        ["Fecha Fin/Obtención", "Fecha Fin", "Fecha de Obtención", "End/obtaining Date", "End date"],
        end_date,
        required=True,
    )

    save_debug(page, "step1_dates_filled")
    submit_and_wait_for_step(
        page,
        "Step 2 (Center)",
        STEP2_ACTIVE_TITLES,
        lambda: click_next(page, required=False),
    )


def step_center(page, data):
    print("Step 2: Center Selection")

    search_terms = build_center_search_terms(data)
    if not search_terms:
        raise RuntimeError("Could not build any center search terms from the classification data.")

    selected = False
    attempted_terms = []
    for term in search_terms:
        attempted_terms.append(term)
        fill_text_field(page, ["Centro", "Center"], term, required=True)
        dismiss_center_autocomplete(page)
        click_if_found(page, ["Buscar", "Search"], timeout_ms=5000)
        wait_for_center_results(page, timeout_ms=7000)
        selected = click_center_result(page, data)
        if selected:
            break

    if not selected:
        save_debug(page, "step2_center_not_selected")
        raise RuntimeError(
            "Could not select a center result. "
            f"Tried search terms: {attempted_terms!r}. "
            f"center_search_text={data.get('center_search_text')!r} "
            f"or institution_name={data.get('institution_name')!r} "
            f"or organizing_entity={data.get('organizing_entity')!r}"
        )

    fill_text_field(
        page,
        ["DescripciÃ³n Adicional", "Descripcion Adicional", "Additional Description"],
        best_center_description_text(data),
    )

    save_debug(page, "step2_center_filled")
    submit_and_wait_for_step(
        page,
        "Step 3 (Name of Merit)",
        STEP3_ACTIVE_TITLES,
        lambda: click_next(page, required=False),
    )


def step_merit_name(page, data):
    print("Step 3: Name of Merit")

    fill_text_field(
        page,
        ["Nombre del Mérito", "Nombre del Merito", "Name of Merit"],
        data.get("merit_name"),
        required=True,
    )

    save_debug(page, "step3_merit_name_filled")
    submit_and_wait_for_step(
        page,
        "Step 4 (Specific values)",
        STEP4_ACTIVE_TITLES,
        lambda: click_next(page, required=False),
    )


def step_specific_values(page, data):
    print("Step 4: Specific values")
    print("Step 4 uses mostly dropdown fields; the automation will normalize Spanish and English labels before selecting.")
    row_records = step4_row_records(page)
    if row_records:
        row_summary = ", ".join(f"{record['row_index']}={record['label']}" for record in row_records)
        print(f"Step 4 rows detected: {row_summary}")

    if data.get("vec_route") == "diplomas_titulos_propios_universitarios":
        scope_value = resolve_scope(data)
        delivery_method = resolve_delivery_method(data)
        training_type = data.get("training_type")
        credits_or_hours = "Horas"
        print(
            "Step 4 dropdowns: "
            f"Ámbito={scope_value or '(blank)'}, "
            f"Metodología={delivery_method or '(blank)'}, "
            f"Tipo={training_type or '(blank)'}, "
            f"Créditos/Horas={credits_or_hours or '(blank)'}"
        )
        step_specific_values_university_diploma(page, data)
        finalize_step4_values(
            page,
            data,
            use_row_layout=bool(row_records),
            training_type=training_type,
            training_type_labels=STEP4_UNIVERSITY_TRAINING_TYPE_LABELS,
            delivery_method=delivery_method,
            scope_value=scope_value,
        )
        save_debug(page, "step4_specific_values_filled")
        submit_and_wait_for_step(
            page,
            "Step 5 (Documents)",
            STEP5_ACTIVE_TITLES,
            lambda: submit_step4(page),
        )
        return

    used_row_layout = fill_standard_specific_values_by_row(page, data)
    if not used_row_layout:
        scope_value = resolve_scope(data)
        if scope_value:
            select_option_field(
                page,
                ["Ámbito", "Ambito", "Scope"],
                scope_value,
                field_kind="scope",
            )
        select_option_field(
            page,
            ["Tipo FormaciÃ³n", "Tipo Formacion", "Training type", "Tipo FormaciÃƒÂ³n"],
            resolve_formacion_continuada_training_type(data) or data.get("training_type"),
            required=True,
            field_kind="training_type",
        )
        select_option_field(
            page,
            ["MetodologÃ­a", "Metodologia", "Delivery method", "MetodologÃƒÂ­a"],
            resolve_delivery_method(data) or data.get("delivery_method"),
            field_kind="delivery_method",
        )

        accredited_choice = resolve_accredited_choice(data)
        if accredited_choice:
            select_option_field(
                page,
                ["Acreditada Recibida", "Accredited received", "Acreditada"],
                accredited_choice,
                field_kind="accredited_choice",
            )

            accredited_received = resolve_accrediting_body(data)
            if accredited_received:
                select_option_field(
                    page,
                    [
                        "Órgano Acreditador Formación Recibida",
                        "Organo Acreditador Formacion Recibida",
                        "Órgano Acreditador",
                        "Organo Acreditador",
                        "Accrediting body",
                    ],
                    accredited_received,
                    field_kind="accredited_received",
                )

        credits_or_hours = resolve_credits_or_hours(data)
        if credits_or_hours:
            select_option_field(
                page,
                ["Créditos/Horas", "Creditos/Horas", "Credits/Hours"],
                credits_or_hours,
                field_kind="credits_or_hours",
            )

    finalize_step4_values(page, data, use_row_layout=used_row_layout)


    save_debug(page, "step4_specific_values_filled")
    submit_and_wait_for_step(
        page,
        "Step 5 (Documents)",
        STEP5_ACTIVE_TITLES,
        lambda: submit_step4(page),
    )


def step_documents(page, data):
    print("Step 5: Document Upload")
    description_candidates = [
        data.get("description"),
        data.get("merit_name"),
        data.get("document_description"),
    ]
    upload_description = ""
    for candidate in description_candidates:
        cleaned = re.sub(r"\s+", " ", str(candidate or "")).strip()
        if not cleaned:
            continue
        if len(cleaned) <= UPLOAD_DESCRIPTION_MAX_LENGTH:
            upload_description = cleaned
            break
        if not upload_description:
            shortened = cleaned[:UPLOAD_DESCRIPTION_MAX_LENGTH].rsplit(" ", 1)[0].strip(" ,;:-")
            upload_description = shortened or cleaned[:UPLOAD_DESCRIPTION_MAX_LENGTH].strip()

    if not upload_description:
        raise RuntimeError("Could not build a valid Step 5 document description.")

    if len(upload_description) > UPLOAD_DESCRIPTION_MAX_LENGTH:
        upload_description = upload_description[:UPLOAD_DESCRIPTION_MAX_LENGTH].strip()

    print(f"Step 5 description: {upload_description!r}")

    fill_text_field(
        page,
        ["DescripciÃ³n", "Descripcion", "Description"],
        upload_description,
        required=True,
    )
    set_checkbox_state(
        page,
        [
            "documentaciÃ³n adjunta sea sÃ³lo el justificante",
            "documentacion adjunta sea solo el justificante",
            "justificante de haber solicitado la certificaciÃ³n del mÃ©rito",
            "justificante de haber solicitado la certificacion del merito",
        ],
        bool(data.get("document_justification_only")),
    )

    upload_pdf_path = build_upload_pdf_path(data["resolved_pdf_path"])
    print(f"Step 5 upload file: {Path(upload_pdf_path).name}")

    uploaded = upload_pdf(page, upload_pdf_path)
    if not uploaded:
        raise RuntimeError("Could not upload the PDF file on the documents step.")

    if not wait_for_document_row(page, timeout_ms=30000):
        save_debug(page, "step5_document_attach_failed")
        upload_error = document_upload_error_message(page)
        if upload_error:
            raise RuntimeError(
                "The PDF file was chosen but Step 5 reported an upload error: "
                f"{upload_error}"
            )
        raise RuntimeError("The PDF file was chosen but did not appear in the Step 5 document table.")

    save_debug(page, "step5_document_uploaded")
    submit_and_wait_for_step(
        page,
        "Step 6 (Signing)",
        STEP6_ACTIVE_TITLES,
        lambda: click_document_continue(page),
    )


def step_sign(page):
    print("Step 6: Signing")

    save_debug(page, "step6_signing_page")

    if not FINAL_SUBMIT:
        print("Stopped on signing page. FINAL_SUBMIT is False.")
        return

    accept_modal_if_present(page)

    if click_if_found_with_retry(page, ["Firmar y registrar", "Sign and submit", "Sign and register"], timeout_ms=6000):
        accept_modal_if_present(page)
        if wait_for_step6_completion(page, timeout_ms=25000):
            save_debug(page, "step6_signed_and_submitted")
            print("Document signed and submitted.")
            return

    signed = click_step6_sign_and_finish(page)
    if not signed:
        save_debug(page, "step6_sign_button_missing")
        raise RuntimeError("Could not find the Step 6 signing button.")

    accept_modal_if_present(page)
    server_signed = click_step6_server_sign(page)
    if not server_signed:
        save_debug(page, "step6_server_sign_button_missing")
        raise RuntimeError("The signing dialog opened, but 'Firmar con Servidor' did not appear.")

    accept_modal_if_present(page)
    if not wait_for_step6_completion(page, timeout_ms=30000):
        save_debug(page, "step6_submit_button_missing")
        raise RuntimeError("Clicked 'Firmar con Servidor', but the Step 6 signing flow did not finish.")

    save_debug(page, "step6_signed_and_submitted")
    print("Document signed and submitted.")


def build_wizard_steps():
    return [
        {
            "name": "Step 1 (Dates)",
            "markers": STEP1_DATE_MARKERS,
            "active_titles": STEP1_ACTIVE_TITLES,
            "field_labels": [
                ["Fecha de Inicio", "Fecha Inicio", "Start date"],
                ["Fecha Fin", "End date"],
            ],
            "runner": step_dates,
        },
        {
            "name": "Step 2 (Center)",
            "markers": STEP2_CENTER_MARKERS,
            "active_titles": STEP2_ACTIVE_TITLES,
            "field_labels": [["Centro", "Center"]],
            "runner": step_center,
        },
        {
            "name": "Step 3 (Name of Merit)",
            "markers": STEP3_MERIT_MARKERS,
            "active_titles": STEP3_ACTIVE_TITLES,
            "field_labels": [["Nombre del MÃ©rito", "Nombre del Merito", "Name of Merit"]],
            "runner": step_merit_name,
        },
        {
            "name": "Step 4 (Specific values)",
            "markers": STEP4_SPECIFIC_MARKERS + STEP4_GENERIC_MARKERS,
            "active_titles": STEP4_ACTIVE_TITLES,
            "field_labels": [
                ["Tipo Formacion", "Tipo FormaciÃ³n", "Training type"],
                ["Creditos/Horas", "CrÃ©ditos/Horas", "Credits/Hours"],
                ["Codigo curso", "CÃ³digo curso", "Course code"],
            ],
            "runner": step_specific_values,
        },
        {
            "name": "Step 5 (Documents)",
            "markers": STEP5_DOCUMENT_MARKERS,
            "active_titles": STEP5_ACTIVE_TITLES,
            "field_labels": [["Descripcion", "DescripciÃ³n", "Description"]],
            "runner": step_documents,
        },
        {
            "name": "Step 6 (Signing)",
            "markers": STEP6_SIGN_MARKERS,
            "active_titles": STEP6_ACTIVE_TITLES,
            "runner": lambda page, data: step_sign(page),
        },
    ]


def detect_current_wizard_step(context, steps, timeout_ms=3000):
    page = attach_to_wizard_page(context, timeout_ms=timeout_ms)
    if not page:
        return None, None

    for index, step in enumerate(steps):
        if page_matches_step(
            page,
            step["markers"],
            active_titles=step["active_titles"],
            field_labels=step.get("field_labels"),
        ):
            return index, page

    return None, page


def wait_for_manual_recovery(context, steps, failed_step_name):
    print()
    print("Inspect or fix the page in that browser window.")
    print("You can leave the wizard on the current step or move it to the next step.")

    while True:
        pages = open_context_pages(context)
        if not pages:
            print("No open pages are available right now. Reopen the VEC wizard in that same browser session, then continue.")
        elif len(pages) == 1 and safe_page_url(pages[0]) == PUBLIC_VEC_URL:
            print("The browser is still on the public VEC landing page. Open the wizard again in that same window, then continue.")

        input("Press Enter when the wizard is ready to continue...")
        step_index, page = detect_current_wizard_step(context, steps, timeout_ms=5000)
        if step_index is not None:
            print(f"Resuming from {steps[step_index]['name']}.")
            return step_index, page

        save_context_debug(context, f"recovery_wait_{failed_step_name}")
        print(f"I still could not detect a wizard step after {failed_step_name}. Keep the wizard visible and try again.")


def page_contains_markers(page, markers):
    for scope in iter_scopes(page):
        if scope_contains_markers(scope, markers):
            return True
    return False


def merit_detail_page_ready(page):
    if not page_contains_markers(page, ["Documentos"]):
        return False

    return page_contains_markers(page, ["Editar", "Eliminar", "Volver", "Detalle"])


def merit_listing_ready(page):
    if merit_detail_page_ready(page):
        return False

    if page_has_wizard_shell(page):
        return False

    for scope in iter_scopes(page):
        candidates = [
            scope.locator("#formMenuNavegacion\\:tipMerTree"),
            scope.locator("[id$='tipMerTree']"),
            scope.locator("#form\\:tabViewMeritos"),
            scope.locator("[id$='tabViewMeritos']"),
            scope.locator("button:has-text('Crear Nuevo')"),
            scope.locator("a:has-text('Crear Nuevo')"),
        ]
        if first_existing(candidates):
            return True

    return page_contains_markers(
        page,
        [
            "BÃºsqueda de MÃ©ritos",
            "Busqueda de Meritos",
            "Listado de MÃ©ritos",
            "Listado de Meritos",
            "Crear Nuevo MÃ©rito",
            "Crear Nuevo Merito",
            "CatÃ¡logo de Tipos de MÃ©ritos",
            "Catalogo de Tipos de Meritos",
        ],
    )


def find_mis_meritos_shortcut(page):
    for scope in iter_scopes(page):
        locator = first_visible(
            [
                scope.get_by_role("button", name="Mis Méritos", exact=False),
                scope.get_by_role("button", name="Mis Meritos", exact=False),
                scope.get_by_role("link", name="Mis Méritos", exact=False),
                scope.get_by_role("link", name="Mis Meritos", exact=False),
                scope.locator("button:has-text('Mis Méritos')"),
                scope.locator("button:has-text('Mis Meritos')"),
                scope.locator("a:has-text('Mis Méritos')"),
                scope.locator("a:has-text('Mis Meritos')"),
                scope.locator("input[value*='Mis M' i]"),
                scope.locator(
                    "xpath=(//*[contains(normalize-space(.), 'Mis Méritos') or contains(normalize-space(.), 'Mis Meritos')]/ancestor-or-self::*[self::a or self::button or @role='button' or contains(@class, 'managementButton')][1])[1]"
                ),
                scope.get_by_text("Mis Méritos", exact=False),
                scope.get_by_text("Mis Meritos", exact=False),
            ]
        )
        if locator:
            return locator
    return None


def dashboard_ready(page):
    if not find_mis_meritos_shortcut(page):
        return False
    return page_contains_markers(page, POST_LOGIN_DASHBOARD_MARKERS)


def click_mis_meritos_shortcut(page):
    locator = find_mis_meritos_shortcut(page)
    if not locator:
        return click_if_found(page, ["Mis MÃ©ritos", "Mis Meritos"], timeout_ms=5000)

    try:
        locator.scroll_into_view_if_needed(timeout=5000)
    except Exception:
        pass

    try:
        locator.click(timeout=5000)
        page.wait_for_timeout(1500)
        return True
    except Exception:
        pass

    try:
        locator.evaluate("(element) => element.click()")
        page.wait_for_timeout(1500)
        return True
    except Exception:
        return click_if_found(page, ["Mis MÃ©ritos", "Mis Meritos"], timeout_ms=5000)


def ensure_mis_meritos_page(page, timeout_ms=15000):
    import time

    if merit_listing_ready(page):
        return page

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        if merit_listing_ready(page):
            return page

        if click_mis_meritos_shortcut(page):
            page.wait_for_timeout(1500)
            if merit_listing_ready(page):
                return page

        if not safe_wait_for_timeout(page, 500):
            break

    return page


def wait_for_logged_in_vec_page(context, timeout_ms=VEC_SESSION_READY_TIMEOUT_MS):
    import time

    end = time.time() + (timeout_ms / 1000)
    last_state = "waiting for QR/access completion"

    while time.time() < end:
        page = attach_to_wizard_page(context, timeout_ms=800) or choose_start_page(context)
        if not page:
            time.sleep(0.5)
            continue

        accept_modal_if_present(page)

        if merit_listing_ready(page):
            return page

        if merit_detail_page_ready(page):
            return page

        if page_contains_markers(page, WIZARD_STEP_MARKERS):
            return page

        if dashboard_ready(page):
            last_state = "logged-in dashboard detected"
            page = ensure_mis_meritos_page(page, timeout_ms=7000)
            if merit_listing_ready(page):
                return page
            last_state = "dashboard detected but 'Mis Méritos' is still opening"
        else:
            last_state = f"current page is {safe_page_title(page)}"

        time.sleep(0.8)

    raise RuntimeError(
        "Timed out while waiting for the logged-in VEC dashboard or merits page. "
        f"Last state: {last_state}"
    )


def return_to_merit_listing(context, timeout_ms=15000):
    import time

    end = time.time() + (timeout_ms / 1000)
    while time.time() < end:
        page = attach_to_wizard_page(context, timeout_ms=1500) or choose_start_page(context)
        if not page:
            return False

        accept_modal_if_present(page)
        page = ensure_mis_meritos_page(page, timeout_ms=5000)
        if merit_listing_ready(page):
            return True

        if merit_detail_page_ready(page):
            if click_if_found_with_retry(page, ["Volver", "Back"], timeout_ms=5000):
                page.wait_for_timeout(1200)
                continue

        if not safe_wait_for_timeout(page, 500):
            break

    page = attach_to_wizard_page(context, timeout_ms=1000) or choose_start_page(context)
    return bool(page and merit_listing_ready(page))


def click_visible_locator(page, locator_list, timeout_ms=5000):
    locator = first_visible(locator_list)
    if not locator:
        return False

    try:
        locator.scroll_into_view_if_needed(timeout=timeout_ms)
    except Exception:
        pass

    try:
        locator.click(timeout=timeout_ms)
        page.wait_for_timeout(1500)
        return True
    except Exception:
        return False


def click_new_merit_button(page):
    if click_if_found(
        page,
        [
            "Crear Nuevo MÃ©rito",
            "Crear Nuevo Merito",
            "Nuevo MÃ©rito",
            "Nuevo Merito",
        ],
        timeout_ms=5000,
    ):
        return True

    for scope in iter_scopes(page):
        direct_locators = [
            scope.locator("button:has-text('Crear Nuevo')"),
            scope.locator("a:has-text('Crear Nuevo')"),
            scope.locator("button:has-text('Nuevo MÃ©rito')"),
            scope.locator("button:has-text('Nuevo Merito')"),
            scope.locator("a:has-text('Nuevo MÃ©rito')"),
            scope.locator("a:has-text('Nuevo Merito')"),
            scope.locator("button.managementButton:has(.ui-icon-plusthick)"),
            scope.locator("a.managementButton:has(.ui-icon-plusthick)"),
            scope.locator("button.managementButton:has(.ui-icon-plus)"),
            scope.locator("a.managementButton:has(.ui-icon-plus)"),
            scope.locator("button[id*='nuevo' i]"),
            scope.locator("a[id*='nuevo' i]"),
            scope.locator("button[id*='crear' i]"),
            scope.locator("a[id*='crear' i]"),
            scope.locator("button[title*='Nuevo' i]"),
            scope.locator("a[title*='Nuevo' i]"),
            scope.locator("button[title*='Crear' i]"),
            scope.locator("a[title*='Crear' i]"),
            scope.locator("button[aria-label*='Nuevo' i]"),
            scope.locator("a[aria-label*='Nuevo' i]"),
            scope.locator("button[aria-label*='Crear' i]"),
            scope.locator("a[aria-label*='Crear' i]"),
            scope.locator("button:has-text('Nuevo')"),
            scope.locator("a:has-text('Nuevo')"),
            scope.locator("button:has-text('Crear')"),
            scope.locator("a:has-text('Crear')"),
        ]
        if click_visible_locator(page, direct_locators, timeout_ms=5000):
            return True

    if click_if_found(page, ["AÃ±adir", "Anadir", "Alta"], timeout_ms=5000):
        return True

    return False


def open_wizard_for_current_document(context, data, timeout_ms=45000):
    import time

    steps = build_wizard_steps()
    end = time.time() + (timeout_ms / 1000)
    last_error = "wizard not detected yet"

    while time.time() < end:
        step_index, page = detect_current_wizard_step(context, steps, timeout_ms=1200)
        if step_index is not None and page:
            return page, step_index

        page = attach_to_wizard_page(context, timeout_ms=1500) or choose_start_page(context)
        if not page:
            last_error = "no browser page is currently available"
            time.sleep(0.5)
            continue

        accept_modal_if_present(page)

        if merit_detail_page_ready(page):
            if click_if_found_with_retry(page, ["Volver", "Back"], timeout_ms=5000):
                accept_modal_if_present(page)
                last_error = "returned from detail page"
                continue

        page = ensure_mis_meritos_page(page, timeout_ms=8000)
        if not merit_listing_ready(page):
            last_error = "could not reach the 'Mis MÃ©ritos' page yet"
            time.sleep(0.6)
            continue

        try:
            page = open_route_from_tree(page, data)
        except Exception as exc:
            last_error = str(exc)
            time.sleep(0.6)
            continue

        if not merit_listing_ready(page):
            last_error = "the merit selection finished, but the script is no longer on the 'Mis Méritos' listing page before 'Crear Nuevo Mérito'"
            time.sleep(0.6)
            continue

        if click_new_merit_button(page):
            accept_modal_if_present(page)
            step_index, wizard_page = detect_current_wizard_step(context, steps, timeout_ms=6000)
            if step_index is not None and wizard_page:
                return wizard_page, step_index
            last_error = "clicked the new merit button, but the wizard did not open"
        else:
            last_error = "the correct merit was selected on the left, but 'Crear Nuevo Mérito' was not found"

        time.sleep(0.6)

    save_context_debug(context, "open_next_wizard_failed")
    raise RuntimeError(
        "Could not auto-open the merit wizard for the next document. "
        f"Last navigation state: {last_error}"
    )


def wait_until_wizard_opens(context, data, batch_number):
    while True:
        try:
            return open_wizard_for_current_document(context, data)
        except Exception as exc:
            print()
            print(f"Automatic wizard opening paused: {exc}")
            save_context_debug(context, f"batch_{batch_number}_open_wizard")
            input(
                "Keep the browser on the VEC dashboard, 'Mis Méritos', the selected merit page, or the wizard page, then press Enter here..."
            )


def run_wizard_steps(context, data, start_index=None):
    steps = build_wizard_steps()
    current_index = 0 if start_index is None else max(0, min(start_index, len(steps) - 1))
    wizard_start_saved = False

    if start_index is None:
        detected_index, _ = detect_current_wizard_step(context, steps, timeout_ms=1500)
        if detected_index is not None:
            current_index = detected_index

    if current_index > 0:
        print(f"Starting wizard automation from {steps[current_index]['name']}.")

    while current_index < len(steps):
        step = steps[current_index]

        try:
            page = require_visible_step(
                context,
                step["name"],
                step["markers"],
                active_titles=step["active_titles"],
                field_labels=step.get("field_labels"),
            )

            if current_index == 0:
                page = ensure_expected_route(page, data)
                page = require_visible_step(
                    context,
                    step["name"],
                    step["markers"],
                    active_titles=step["active_titles"],
                    field_labels=step.get("field_labels"),
                )
                if not wizard_start_saved:
                    print("Attached to the wizard page.")
                    save_debug(page, "wizard_start")
                    wizard_start_saved = True

            step["runner"](page, data)
            current_index += 1
        except Exception as exc:
            print()
            print(f"Automation paused on {step['name']}: {exc}")
            save_context_debug(context, f"paused_{current_index + 1}")
            recovered_index, _ = wait_for_manual_recovery(context, steps, step["name"])
            current_index = recovered_index

    print("Wizard automation finished.")


def main():
    args = parse_args()
    validate_route_tree_configuration()
    target_successes = max(1, int(args.target_successes or DEFAULT_TARGET_SUCCESSFUL_UPLOADS))
    classification_backend = str(args.classification_backend or DEFAULT_CLASSIFICATION_BACKEND).strip().lower() or DEFAULT_CLASSIFICATION_BACKEND

    print(f"Script version: {SCRIPT_VERSION}")
    print(f"Final submit enabled: {FINAL_SUBMIT}")
    print(f"Classification backend: {classification_backend}")
    print(f"Target successful uploads: {target_successes}")
    print()
    print("Flow:")
    print("1. Script preloads the next complete PDF classification")
    print("2. Browser opens")
    print("3. Complete QR/access in that same browser window")
    print("4. Script detects the VEC dashboard, opens 'Mis Méritos', chooses the merit, fills the wizard, and repeats")
    print()
    print("Preloading the first complete PDF classification before opening the VEC session...")
    next_batch_payload = classify_next_ready_document(
        classification_backend=classification_backend,
    )

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            USER_DATA_DIR,
            headless=False,
            viewport={"width": 1400, "height": 900},
            accept_downloads=True,
        )

        page = open_public_vec_start_page(context)
        print(f"Startup page ready: {safe_page_url(page)}")

        print("Waiting for QR/access completion so the script can continue from the VEC dashboard automatically...")
        try:
            page = wait_for_logged_in_vec_page(context)
            print(f"Logged-in VEC page detected: {safe_page_url(page)}")
        except Exception as exc:
            print()
            print(f"Automatic post-login detection paused: {exc}")
            input(
                "When the VEC dashboard, 'Mis Méritos' page, or the wizard page is visible in that browser window, press Enter here..."
            )
            page = wait_for_logged_in_vec_page(context, timeout_ms=30000)
        pending_sync_result = flush_pending_excel_syncs()
        if pending_sync_result["synced"]:
            print(f"Synchronized {pending_sync_result['synced']} queued Excel update(s).")
        if pending_sync_result["remaining"]:
            print(
                "Warning: some Excel updates are still queued because the workbook is not writable right now "
                f"({pending_sync_result['remaining']} pending)."
            )

        successful_uploads = 0
        while successful_uploads < target_successes:
            batch_number = successful_uploads + 1
            print()
            print(f"=== Batch document {batch_number}/{target_successes} ===")

            data, _, missing_required, missing_recommended, bad_dates = next_batch_payload
            blockers = classification_blockers(data, missing_required, missing_recommended, bad_dates)
            if blockers:
                raise RuntimeError(
                    "The batch loop selected a blocked classification unexpectedly: "
                    + " | ".join(blockers)
                )

            try:
                _, start_index = wait_until_wizard_opens(context, data, batch_number)
            except Exception as exc:
                print()
                print(f"Automatic wizard opening paused: {exc}")
                save_context_debug(context, f"batch_{batch_number}_open_wizard")
                input(
                    "Open the VEC dashboard, 'Mis Méritos', or the next wizard page in that same browser window, then press Enter here..."
                )
                _, start_index = open_wizard_for_current_document(context, data)

            run_wizard_steps(context, data, start_index=start_index)
            archived_pdf_path = ""
            if FINAL_SUBMIT:
                try:
                    archived_pdf_path = archive_uploaded_pdf(data)
                    if archived_pdf_path:
                        print(f"Archived uploaded PDF to: {archived_pdf_path}")
                except Exception as exc:
                    print(f"Warning: the uploaded PDF could not be archived automatically: {exc}")
                excel_sync_result = sync_successful_upload_to_excel(data, archived_pdf_path=archived_pdf_path)
                if excel_sync_result["status"] == "synced":
                    print(
                        "Updated the tracking workbook after successful submission "
                        f"(row {excel_sync_result['excel_row']})."
                    )
                elif excel_sync_result["status"] == "logged_only":
                    print("Logged the successful submission in the workbook audit sheet.")
                elif excel_sync_result["status"] == "already_recorded":
                    print("This successful submission was already present in the tracking workbook.")
                else:
                    print(
                        "Warning: the workbook update was queued for retry because it could not be saved right now: "
                        f"{excel_sync_result['error']}"
                    )

                pending_sync_result = flush_pending_excel_syncs()
                if pending_sync_result["synced"]:
                    print(f"Synchronized {pending_sync_result['synced']} queued Excel update(s).")
                if pending_sync_result["remaining"]:
                    print(f"Excel updates still queued: {pending_sync_result['remaining']}")

            successful_uploads += 1
            print(f"Successful uploads completed: {successful_uploads}/{target_successes}")

            if FINAL_SUBMIT:
                if return_to_merit_listing(context, timeout_ms=20000):
                    print("Returned to the merit listing page for the next 'Crear Nuevo MÃ©rito'.")
                else:
                    print("Warning: could not auto-return to the merit listing page yet.")

            if successful_uploads < target_successes:
                print()
                print("Preloading the next complete PDF classification for the following merit...")
                next_batch_payload = classify_next_ready_document(
                    classification_backend=classification_backend,
                )

        input("Press Enter to close the browser...")
        safe_close_context(context)


if __name__ == "__main__":
    main()
